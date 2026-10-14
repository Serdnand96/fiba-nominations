"""Helpers chicos para exportar planillas .xlsx con openpyxl.

Antes cada export repetía su propio `_autosize` / `_save` y armaba el
`Response` a mano (logistics_import.py, budget_import.py + budget.py). Esto es
el mismo código, una sola vez, para los exports nuevos (pagos, gastos,
resumen de budget). Los viejos no se migraron a propósito: funcionan y moverlos
no aporta nada.

Uso típico:

    wb = Workbook(); ws = wb.active
    ws.append(headers); style_header(ws)
    for row in rows: ws.append([safe_text(v) for v in row])
    autosize(ws)
    return xlsx_response(workbook_bytes(wb), "pagos.xlsx")
"""
import io
import re

from fastapi.responses import Response
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

# Una celda de texto que arranca con alguno de estos la interpreta Excel como
# fórmula (CSV/formula injection). Mismo criterio que `_safe` en budget_import.
_FORMULA_LEAD = ("=", "+", "-", "@")
_SAFE_FILENAME_RE = re.compile(r"[^\w\-\.]")


def safe_text(value):
    """Texto que ninguna planilla vaya a interpretar como fórmula.

    Solo toca strings: los números (montos) tienen que seguir siendo números
    para que la planilla sume. Un monto negativo nunca llega como string.
    """
    if isinstance(value, str) and value.startswith(_FORMULA_LEAD):
        return "'" + value
    return value


def style_header(ws, row: int = 1) -> None:
    """Encabezado navy del DS, en negrita, con la fila congelada debajo."""
    for cell in ws[row]:
        cell.font = Font(bold=True, color="FFFFFFFF")
        cell.fill = PatternFill("solid", fgColor="FF0C2340")   # navy del DS
        cell.alignment = Alignment(vertical="center")
    # String y no ws.cell(...).coordinate: ws.cell() CREA la celda, y con eso el
    # próximo ws.append() arranca una fila más abajo (fila en blanco).
    ws.freeze_panes = f"A{row + 1}"


def autosize(ws, min_width: int = 10, max_width: int = 40) -> None:
    """Ancho de columna según el contenido más largo, con tope.

    Copiado de `_autosize` (logistics_import.py): sin el tope, una columna de
    comentarios largos empuja todo lo demás fuera de la pantalla.
    """
    for column in ws.columns:
        width = max((len(str(c.value)) for c in column if c.value is not None), default=0)
        ws.column_dimensions[column[0].column_letter].width = min(max(width + 2, min_width), max_width)


def money_format(ws, columns, min_row: int = 2, fmt: str = "#,##0.00") -> None:
    """Formato de moneda para las columnas (1-based) indicadas."""
    for col in columns:
        for row in ws.iter_rows(min_row=min_row, min_col=col, max_col=col):
            for cell in row:
                cell.number_format = fmt


def workbook_bytes(wb: Workbook) -> bytes:
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def xlsx_response(content: bytes, filename: str) -> Response:
    """Response de descarga, como GET /budget/lines/export.xlsx.

    `no-store` porque estos exports llevan montos: que no quede una copia en
    un proxy o en el cache del navegador de una máquina compartida.
    """
    name = _SAFE_FILENAME_RE.sub("_", filename or "export.xlsx") or "export.xlsx"
    if not name.lower().endswith(".xlsx"):
        name += ".xlsx"
    return Response(
        content=content,
        media_type=XLSX_MEDIA_TYPE,
        headers={
            "Content-Disposition": f'attachment; filename="{name}"',
            "Cache-Control": "private, no-store",
        },
    )
