"""Budget — gastos, ingresos y catálogos por departamento.

Complementa `payments`, que solo sabe pagar a una persona nominada a un evento
(`payments.nomination_id` es NOT NULL UNIQUE). Acá entran los dos casos que ese
modelo no cubre:

  1. gasto de departamento sin evento  → licencias, internet, leasing
  2. gasto de evento sin persona       → shipping, branding, seguros

Cada gasto lleva tres dimensiones: **departamento** (quién gasta), **cuenta**
(en qué) y **competencia** (opcional; NULL = gasto general). Departamento y
cuenta son dimensiones separadas a propósito: una competencia cruza
departamentos — la línea "TV Production / Streaming" del presupuesto de
Competitions 2027 es plata de Comms.

Solo `status = 'paid'` consume presupuesto; `approved` se reporta aparte como
comprometido.

Datos financieros sensibles (montos, datos bancarios de proveedores, facturas):
las tablas son backend-only (RLS on, sin políticas) y el router va detrás del
permiso `budget`, view separado de edit.

Ver BUDGET_MODULE.md para el diseño completo.
"""
import json
import logging
import os
import re
import uuid
from datetime import date, datetime, timezone
from typing import Optional

import httpx
from fastapi import APIRouter, Depends, File, Form, HTTPException, Query, Request, UploadFile
from fastapi.responses import Response

from api._lib.auth import has_edit, is_superadmin_request, require_edit, require_superadmin, require_view
from api._lib.budget_accounts import ACCOUNT_BY_ROLE, TRAVEL_ACCOUNT_BY_ROLE, travel_account_for_role
from api._lib.database import supabase
from api._lib.services.finance_history import history_for, record_transition
from api._lib.services.notify import budget_approvers_for, notify
from api._lib.schemas import (
    AccountCreate,
    AccountRemap,
    AccountUpdate,
    AssumptionsPut,
    BudgetAccessPut,
    BudgetLineCreate,
    BudgetLineUpdate,
    BudgetProject,
    ExpenseApprove,
    ExpenseCreate,
    ExpenseTransition,
    ExpenseUpdate,
    FeeScheduleApply,
    FeeScheduleUpdate,
    HeadcountPut,
    RecurringCreate,
    RecurringGenerate,
    RecurringUpdate,
    RevenueCreate,
    RevenueUpdate,
    VendorCreate,
    VendorUpdate,
)

logger = logging.getLogger(__name__)

router = APIRouter(
    prefix="/budget",
    tags=["budget"],
    dependencies=[Depends(require_view("budget"))],
)

_BUCKET = "nominations"
_ATTACH_PREFIX = "expenses"
_MAX_UPLOAD_BYTES = 10 * 1024 * 1024  # 10 MB
_ALLOWED_CONTENT = ("application/pdf", "image/")
_SAFE_FILENAME_RE = re.compile(r"[^\w\s\-\.\(\)]")

_EXPENSE_STATUSES = {"draft", "approved", "paid", "rejected", "cancelled"}
_REVENUE_STATUSES = {"expected", "received", "cancelled"}
_PAYEE_TYPES = {"vendor", "employee", "other"}
_FREQUENCIES = {"monthly", "quarterly", "annual"}
_PERIOD_RE = re.compile(r"^\d{4}-(0[1-9]|1[0-2])$")

# Campos que un PATCH puede cambiar pero nunca vaciar. Los updates usan
# model_dump(exclude_unset=True) — no `if v is not None` — para distinguir "no
# lo mandé" de "ponelo en null"; sin eso no habría forma de sacarle la
# competencia a un gasto y pasarlo a general. El precio es tener que rechazar a
# mano los null sobre las columnas NOT NULL.
_REQUIRED_ON_UPDATE = {
    "expenses": {"department_code", "account_code", "description", "amount", "expense_date", "status", "payee_type"},
    "revenues": {"department_code", "account_code", "source_name", "amount", "status"},
    "recurring_expenses": {"department_code", "account_code", "description", "amount", "frequency", "start_date"},
    "budget_lines": {"department_code", "account_code", "description", "amount", "kind", "escalation_pct", "source"},
    "vendors": {"name"},
    "accounts": {"label", "kind"},
}


def _patch_updates(data, table: str) -> dict:
    """Campos efectivamente enviados, rechazando null sobre columnas NOT NULL."""
    updates = data.model_dump(exclude_unset=True)
    if not updates:
        raise HTTPException(status_code=400, detail="Nothing to update")
    for field in _REQUIRED_ON_UPDATE.get(table, set()) & updates.keys():
        if updates[field] is None:
            raise HTTPException(status_code=400, detail=f"{field} cannot be null")
    return updates


# ─── Helpers ─────────────────────────────────────────────────────────────────
def _user_id(request: Request) -> Optional[str]:
    user = getattr(request.state, "user", None)
    if isinstance(user, dict):
        return user.get("id")
    return None


def _now_iso() -> str:
    return datetime.now(timezone.utc).isoformat()


def _valid_department(code: str) -> bool:
    rows = supabase.table("departments").select("code").eq("code", code).eq("active", True).execute().data
    return bool(rows)


def _account(code: str) -> Optional[dict]:
    rows = supabase.table("accounts").select("code, kind, active").eq("code", code).execute().data
    return rows[0] if rows else None


def _check_account(code: str, expected_kind: str) -> None:
    """La cuenta existe, está activa y es del tipo correcto.

    El CHECK de Postgres no puede mirar `accounts.kind` desde `expenses`, así
    que la coherencia gasto↔cuenta de gasto se valida acá. Sin esto se podría
    imputar un gasto a una cuenta de ingreso y el reporte quedaría torcido sin
    que nada falle.
    """
    acc = _account(code)
    if not acc or not acc.get("active"):
        raise HTTPException(status_code=400, detail="Unknown or inactive account")
    if acc.get("kind") != expected_kind:
        raise HTTPException(
            status_code=400,
            detail=f"Account {code} is a {acc.get('kind')} account, expected {expected_kind}",
        )


def _check_payee(payee_type: str, vendor_id, employee_id, payee_name) -> None:
    """Refleja el CHECK expense_payee_present, para dar un 400 con mensaje.

    Sin esto el constraint dispara un 500 genérico de Postgres y el usuario no
    se entera de qué le falta.
    """
    if payee_type not in _PAYEE_TYPES:
        raise HTTPException(status_code=400, detail="Invalid payee_type")
    if payee_type == "vendor" and not vendor_id:
        raise HTTPException(status_code=400, detail="A vendor is required when payee_type is 'vendor'")
    if payee_type == "employee" and not employee_id:
        raise HTTPException(status_code=400, detail="An employee is required when payee_type is 'employee'")
    if payee_type == "other" and not (payee_name or "").strip():
        raise HTTPException(status_code=400, detail="payee_name is required when payee_type is 'other'")


def _check_paid_coherence(status: str, payment_date) -> None:
    """Refleja el CHECK expense_paid_has_date."""
    if status == "paid" and not payment_date:
        raise HTTPException(status_code=400, detail="A payment date is required to mark an expense as paid")
    if status != "paid" and payment_date:
        raise HTTPException(status_code=400, detail="Only a paid expense can carry a payment date")


# ─── Scope por departamento ──────────────────────────────────────────────────
#
# PRIMER filtrado por fila del sistema. El resto de los módulos abre o cierra un
# endpoint entero; acá el designado de IT tiene que ver los gastos de IT y no
# los de Competitions.
#
# RLS no sirve: el backend pega con service_role, que la bypassa por diseño. El
# recorte vive acá y hay que aplicarlo en CADA query. Una lectura de budget sin
# _scoped() es un agujero P0, igual que un endpoint sin require_view.
#
# Falla cerrado: sin filas en budget_access el usuario no ve nada, aunque tenga
# el permiso del módulo. Es deliberado — el default opuesto convertiría un
# olvido del admin en una fuga de datos financieros.
def _access_rows(request: Request) -> list:
    cached = getattr(request.state, "_budget_access", None)
    if cached is not None:
        return cached
    uid = _user_id(request)
    rows = []
    if uid:
        rows = supabase.table("budget_access").select("*").eq("user_id", uid).execute().data or []
    request.state._budget_access = rows
    return rows


_ACCESS_COLUMN = {"view": "can_view", "edit": "can_edit", "approve": "can_approve"}


def _departments_for(request: Request, action: str) -> Optional[list]:
    """Departamentos permitidos, o None si son todos.

    action: 'view' | 'edit' | 'approve'. 'approve' (migración 048) es una
    columna aparte y NO se deduce de can_edit: el que carga un gasto y el que lo
    aprueba son roles distintos. Si la 048 no está aplicada, las filas no traen
    `can_approve` y el resultado es [] — falla cerrado, nadie aprueba salvo el
    superadmin.

    Una acción desconocida levanta: antes caía en can_view, y un typo en el
    nombre de la acción se convertía en permiso de lectura en silencio.
    """
    if is_superadmin_request(request):
        return None
    col = _ACCESS_COLUMN.get(action)
    if col is None:
        raise ValueError(f"Unknown budget access action {action!r}")
    rows = _access_rows(request)

    def grants(r: dict) -> bool:
        # Aprobar lo que no se ve no tiene sentido; el PUT ya lo impide, esto
        # es por si alguien carga la fila a mano en SQL.
        return bool(r.get(col)) and (action != "approve" or bool(r.get("can_view")))

    if any(r.get("department_code") == "*" and grants(r) for r in rows):
        return None
    return sorted({r["department_code"] for r in rows if grants(r) and r["department_code"] != "*"})


def _allowed_departments(request: Request, action: str = "view") -> Optional[list]:
    """Departamentos que una LECTURA puede devolver para `action`.

    'view' es el scope normal. Cualquier otra acción se intersecta con 'view':
    la bandeja "pendientes de aprobación" muestra drafts de los departamentos
    donde el caller aprueba Y ve. None = todos.
    """
    viewable = _departments_for(request, "view")
    if action == "view":
        return viewable
    other = _departments_for(request, action)
    if viewable is None:
        return other
    if other is None:
        return viewable
    return sorted(set(viewable) & set(other))


def _scoped(q, request: Request, requested: Optional[str] = None, action: str = "view"):
    """Recorte por departamento, combinado con el filtro que pidió el usuario.

    ⚠️ Los dos se resuelven ACÁ y se aplica UN SOLO filtro, a propósito: el
    query builder guarda un filtro por columna (`self._params[column] = ...`),
    así que un `.eq("department_code", x)` encadenado después de un `.in_()`
    PISA el scope silenciosamente. Encadenarlos sería un bypass: bastaría pedir
    `?department=finance` para leer un departamento ajeno.

    Va en TODA lectura del módulo. Falla cerrado: lista vacía → PostgREST
    `in.()` → no matchea nada.

    `action='approve'` recorta además a los departamentos donde el caller
    aprueba (bandeja de pendientes). La intersección con lo que pidió el
    usuario se resuelve en Python, antes de armar el único filtro.
    """
    allowed = _allowed_departments(request, action)
    if requested:
        if allowed is not None and requested not in allowed:
            return q.in_("department_code", [])   # fuera de scope → sin resultados
        return q.eq("department_code", requested)
    if allowed is None:
        return q
    return q.in_("department_code", allowed)


def _can_view_department(request: Request, code: Optional[str]) -> bool:
    depts = _departments_for(request, "view")
    return depts is None or (code in depts)


def _assert_can_edit(request: Request, code: Optional[str]) -> None:
    depts = _departments_for(request, "edit")
    if depts is None:
        return
    if not code or code not in depts:
        raise HTTPException(status_code=403, detail=f"No edit access to department '{code}'")


def _assert_can_approve(request: Request, code: Optional[str]) -> None:
    """403 si el caller no aprueba gastos de `code` (superadmin siempre pasa)."""
    depts = _departments_for(request, "approve")
    if depts is None:
        return
    if not code or code not in depts:
        raise HTTPException(status_code=403, detail=f"No approval rights on department '{code}'")


def _is_general_admin(request: Request) -> bool:
    """Superadmin o fila `*` con can_edit en budget_access."""
    return _departments_for(request, "edit") is None


def _assert_general_admin(request: Request) -> None:
    """Solo el admin general del módulo (superadmin o `*` con can_edit).

    Para lo que no pertenece a ningún departamento y por lo tanto no se puede
    recortar por fila: el plan de cuentas y los datos bancarios / fiscales de
    los proveedores. Con `require_edit("budget")` solo, el designado de IT podía
    renombrar una cuenta de Competitions o cambiar el CBU de un proveedor ajeno
    — y eso último es el fraude de manual (redirigir el pago a otra cuenta).
    """
    if not _is_general_admin(request):
        raise HTTPException(status_code=403, detail="Only a general budget admin can do this")


def _row_or_403(table: str, row_id: str, request: Request, action: str = "view") -> dict:
    """Trae una fila y verifica el acceso a su departamento.

    Devuelve 404 tanto si no existe como si el caller no la puede ver: un 403
    distinguible confirmaría que el registro existe.
    """
    rows = supabase.table(table).select("*").eq("id", row_id).execute().data
    if not rows:
        raise HTTPException(status_code=404, detail="Not found")
    row = rows[0]
    if not _can_view_department(request, row.get("department_code")):
        raise HTTPException(status_code=404, detail="Not found")
    if action == "edit":
        _assert_can_edit(request, row.get("department_code"))
    return row


def _extract_storage_key(storage_path: Optional[str]) -> Optional[str]:
    """Object key dentro del bucket `nominations`, en los 3 formatos soportados."""
    if not storage_path:
        return None
    if storage_path.startswith(f"storage://{_BUCKET}/"):
        return storage_path[len(f"storage://{_BUCKET}/"):]
    if f"/storage/v1/object/public/{_BUCKET}/" in storage_path:
        return storage_path.split(f"/storage/v1/object/public/{_BUCKET}/", 1)[1]
    if f"/storage/v1/object/{_BUCKET}/" in storage_path:
        return storage_path.split(f"/storage/v1/object/{_BUCKET}/", 1)[1]
    return None


def _delete_from_storage(storage_path: Optional[str]) -> None:
    """Borrado best-effort de un adjunto."""
    key = _extract_storage_key(storage_path)
    if not key:
        return
    try:
        supabase.storage.from_(_BUCKET).remove([key])
    except Exception as e:
        logger.warning(f"[storage cleanup] could not remove {key}: {e}")


# ─── Administración del acceso por departamento ──────────────────────────────
# Solo superadmin: quién ve qué plata no lo decide alguien que ya tiene el
# módulo — si no, un designado podría ampliarse el propio alcance.
@router.get("/access/{user_id}", dependencies=[Depends(require_superadmin)])
def get_budget_access(user_id: str):
    return (
        supabase.table("budget_access")
        .select("*")
        .eq("user_id", user_id)
        .execute()
        .data
    ) or []


@router.put("/access/{user_id}", dependencies=[Depends(require_superadmin)])
def put_budget_access(user_id: str, data: BudgetAccessPut):
    """Reemplaza el acceso del usuario. Sin ítems = sin acceso a ningún dato."""
    valid = {d["code"] for d in (supabase.table("departments").select("code").execute().data or [])}
    valid.add("*")
    for item in data.items:
        if item.department_code not in valid:
            raise HTTPException(status_code=400, detail=f"Unknown department '{item.department_code}'")
        # can_edit sin can_view no significa nada y dejaría una fila que el
        # filtro de lectura ignora pero el de escritura acepta.
        if item.can_edit and not item.can_view:
            raise HTTPException(status_code=400, detail="can_edit requires can_view")
        # Mismo criterio para aprobar: no se aprueba lo que no se ve. can_approve
        # NO exige can_edit — el aprobador puede ser alguien que no carga.
        if item.can_approve and not item.can_view:
            raise HTTPException(status_code=400, detail="can_approve requires can_view")

    supabase.table("budget_access").delete().eq("user_id", user_id).execute()
    kept = [i for i in data.items if i.can_view or i.can_edit or i.can_approve]
    if kept:
        supabase.table("budget_access").insert([{
            "user_id": user_id,
            "department_code": i.department_code,
            "can_view": i.can_view,
            "can_edit": i.can_edit,
            "can_approve": i.can_approve,
        } for i in kept]).execute()
    return {"ok": True, "count": len(kept)}


@router.get("/access/me/scope")
def my_scope(request: Request):
    """Qué ve y qué edita el caller. Lo usa el frontend para armar la UI."""
    return {
        "is_superadmin": is_superadmin_request(request),
        "view": _departments_for(request, "view"),   # None = todos
        "edit": _departments_for(request, "edit"),
        # Departamentos donde aprueba/rechaza gastos (migración 048). Ojo: aun
        # con el departamento acá, nadie aprueba un gasto que cargó él mismo
        # (el superadmin sí, con motivo) — el front lo resuelve por fila.
        "approve": _departments_for(request, "approve"),
    }


# ─── Catálogos ───────────────────────────────────────────────────────────────
@router.get("/departments")
def list_departments(request: Request, include_inactive: bool = Query(False)):
    """Solo los departamentos que el caller puede ver.

    Recortado además de por prolijidad: los selects del frontend se arman con
    esto, y ofrecer un departamento que el backend después rechaza es una
    trampa para el usuario.
    """
    q = supabase.table("departments").select("*").order("sort")
    if not include_inactive:
        q = q.eq("active", True)
    rows = q.execute().data or []
    allowed = _departments_for(request, "view")
    if allowed is not None:
        rows = [r for r in rows if r["code"] in allowed]
    editable = _departments_for(request, "edit")
    for r in rows:
        r["can_edit"] = editable is None or r["code"] in editable
    return rows


@router.get("/accounts")
def list_accounts(request: Request, kind: Optional[str] = Query(None),
                  include_inactive: bool = Query(False)):
    """El plan de cuentas, con qué departamentos usa cada cuenta.

    `used_by` no es una columna: la cuenta **no pertenece** a un departamento
    (decisión #2 — son dimensiones separadas, y "TV Production" es plata de
    Comms dentro del presupuesto de Competitions). Se deriva de lo que hay
    cargado para que la UI pueda ordenar el desplegable por el departamento en
    el que estás parado: con IT elegido, las 56 cuentas de Competitions y Comms
    son ruido.

    Va recortado por el scope del caller: solo aparecen los departamentos que
    puede ver, así el listado no filtra en qué gasta un área ajena.
    """
    q = supabase.table("accounts").select("*").order("sort")
    if kind:
        if kind not in ("expense", "revenue"):
            raise HTTPException(status_code=400, detail="kind must be 'expense' or 'revenue'")
        q = q.eq("kind", kind)
    if not include_inactive:
        q = q.eq("active", True)
    accounts = q.execute().data or []

    used = {}
    for table in ("budget_lines", "expenses", "revenues", "recurring_expenses"):
        rows = _scoped(
            supabase.table(table).select("account_code, department_code"), request,
        ).execute().data or []
        for row in rows:
            if row.get("account_code"):
                used.setdefault(row["account_code"], set()).add(row["department_code"])
    for account in accounts:
        account["used_by"] = sorted(used.get(account["code"], set()))
    return accounts


@router.post("/accounts", status_code=201, dependencies=[Depends(require_edit("budget"))])
def create_account(data: AccountCreate, request: Request):
    # El plan de cuentas es transversal a los departamentos: no hay fila que
    # recortar, así que lo administra solo el admin general.
    _assert_general_admin(request)
    if data.kind not in ("expense", "revenue"):
        raise HTTPException(status_code=400, detail="kind must be 'expense' or 'revenue'")
    code = (data.code or "").strip()
    if not code:
        raise HTTPException(status_code=400, detail="code is required")
    if _account(code):
        raise HTTPException(status_code=409, detail="An account with that code already exists")
    record = {k: v for k, v in data.model_dump().items() if v is not None}
    record["code"] = code
    return supabase.table("accounts").insert(record).execute().data[0]


@router.patch("/accounts/{code}", dependencies=[Depends(require_edit("budget"))])
def update_account(code: str, data: AccountUpdate, request: Request):
    _assert_general_admin(request)
    updates = _patch_updates(data, "accounts")
    if "kind" in updates and updates["kind"] not in ("expense", "revenue"):
        raise HTTPException(status_code=400, detail="kind must be 'expense' or 'revenue'")
    result = supabase.table("accounts").update(updates).eq("code", code).execute()
    if not result.data:
        raise HTTPException(status_code=404, detail="Account not found")
    return result.data[0]


def _roles_mapped_to(code: str) -> list:
    """Roles cuyo pago imputa a `code` según api/_lib/budget_accounts.py."""
    return (
        [f"{role} (fees)" for role, acc in ACCOUNT_BY_ROLE.items() if acc == code]
        + [f"{role} (travel)" for role, acc in TRAVEL_ACCOUNT_BY_ROLE.items() if acc == code]
    )


@router.post("/accounts/{code}/remap", dependencies=[Depends(require_edit("budget"))])
def remap_account(code: str, data: AccountRemap, request: Request):
    """Le asigna el código oficial de Finance a una cuenta (típicamente COMP-*).

    Es un UPDATE de la PK: todas las FK hacia `accounts.code` (expenses,
    revenues, recurring_expenses, budget_lines, payments, fee_schedule y el
    propio parent_code) son ON UPDATE CASCADE, así que los gastos, líneas y
    pagos ya cargados se mueven solos. Se baja `pending_mapping` porque el
    código deja de ser provisorio.

    Lo único que el CASCADE no alcanza es código: `budget_accounts.py` tiene los
    COMP-* de fees/travel por rol hardcodeados. Remapear una de esas cuentas sin
    tocar ese archivo deja al próximo pago de ese rol escribiendo un código que
    ya no existe, y el FK lo rechaza. Por eso 409 salvo `force` — y `force` es
    para cuando el cambio de budget_accounts.py ya está en el mismo deploy.

    Solo admin general: el plan de cuentas es transversal a los departamentos.
    """
    _assert_general_admin(request)
    new_code = (data.new_code or "").strip()
    if not new_code:
        raise HTTPException(status_code=400, detail="new_code is required")
    if new_code == code:
        raise HTTPException(status_code=400, detail="new_code is the same as the current code")
    if not _account(code):
        raise HTTPException(status_code=404, detail="Account not found")
    if _account(new_code):
        raise HTTPException(status_code=409, detail="An account with that code already exists")

    roles = _roles_mapped_to(code)
    if roles and not data.force:
        raise HTTPException(
            status_code=409,
            detail=(
                f"Account {code} is used to allocate payments for {', '.join(roles)}. "
                f"actualizar api/_lib/budget_accounts.py ({code} → {new_code}) in the same "
                "deploy and retry with force=true"
            ),
        )
    if roles:
        logger.warning(
            "[budget] account %s remapped to %s with force; budget_accounts.py maps it for %s",
            code, new_code, ", ".join(roles),
        )

    result = (
        supabase.table("accounts")
        .update({"code": new_code, "pending_mapping": False})
        .eq("code", code)
        .execute()
    )
    if not result.data:
        raise HTTPException(status_code=404, detail="Account not found")
    return result.data[0]


# Datos bancarios y fiscales del proveedor: los ve y los escribe solo el admin
# general. Un proveedor no es de un departamento (Adobe le factura a IT y a
# Comms), así que no se puede recortar por fila como un gasto; y cambiar el
# `bank_info` es exactamente cómo se desvía un pago. El resto ve
# `has_bank_info` para saber si falta cargarlo, sin el dato.
_VENDOR_SENSITIVE = ("bank_info", "tax_id")


def _shape_vendor(row: dict, admin: bool) -> dict:
    row["has_bank_info"] = bool((row.get("bank_info") or "").strip())
    if not admin:
        for field in _VENDOR_SENSITIVE:
            row.pop(field, None)
    return row


@router.get("/vendors")
def list_vendors(request: Request, search: Optional[str] = Query(None),
                 include_inactive: bool = Query(False)):
    admin = _is_general_admin(request)
    q = supabase.table("vendors").select("*").order("name")
    if not include_inactive:
        q = q.eq("active", True)
    rows = q.execute().data or []
    if search:
        s = search.lower()
        # Buscar por tax_id solo si lo puede ver: si no, la búsqueda sirve de
        # oráculo para adivinarlo de a un carácter.
        rows = [
            r for r in rows
            if s in (r.get("name") or "").lower()
            or (admin and s in (r.get("tax_id") or "").lower())
            or s in (r.get("country") or "").lower()
        ]
    return [_shape_vendor(r, admin) for r in rows]


@router.post("/vendors", status_code=201, dependencies=[Depends(require_edit("budget"))])
def create_vendor(data: VendorCreate, request: Request):
    admin = _is_general_admin(request)
    name = (data.name or "").strip()
    if not name:
        raise HTTPException(status_code=400, detail="name is required")
    record = {k: v for k, v in data.model_dump().items() if v is not None}
    if not admin and any(f in record for f in _VENDOR_SENSITIVE):
        raise HTTPException(status_code=403, detail="Only a general budget admin can set bank or tax data")
    record["name"] = name
    record["created_by"] = _user_id(request)
    try:
        row = supabase.table("vendors").insert(record).execute().data[0]
    except Exception as e:
        # idx_vendors_name_unique es case-insensitive: "Adobe" y "adobe" chocan.
        if "idx_vendors_name_unique" in str(e):
            raise HTTPException(status_code=409, detail="A vendor with that name already exists")
        raise
    return _shape_vendor(row, admin)


@router.patch("/vendors/{vendor_id}", dependencies=[Depends(require_edit("budget"))])
def update_vendor(vendor_id: str, data: VendorUpdate, request: Request):
    admin = _is_general_admin(request)
    updates = _patch_updates(data, "vendors")
    # exclude_unset: mandar el campo en null también es escribirlo (borrarlo).
    if not admin and any(f in updates for f in _VENDOR_SENSITIVE):
        raise HTTPException(status_code=403, detail="Only a general budget admin can change bank or tax data")
    updates["updated_at"] = _now_iso()
    try:
        result = supabase.table("vendors").update(updates).eq("id", vendor_id).execute()
    except Exception as e:
        if "idx_vendors_name_unique" in str(e):
            raise HTTPException(status_code=409, detail="A vendor with that name already exists")
        raise
    if not result.data:
        raise HTTPException(status_code=404, detail="Vendor not found")
    # La respuesta trae la fila entera: sin recortarla, un PATCH del nombre le
    # devolvía al designado el bank_info que el GET le esconde.
    return _shape_vendor(result.data[0], admin)


@router.delete("/vendors/{vendor_id}", dependencies=[Depends(require_edit("budget"))])
def delete_vendor(vendor_id: str, request: Request):
    """Baja lógica si el proveedor tiene gastos; borrado real si nunca se usó.

    Los FK son ON DELETE SET NULL, así que un DELETE real no rompería nada —
    pero dejaría gastos históricos sin proveedor y sin forma de recuperarlo.

    Solo admin general: el catálogo es compartido entre departamentos y dar de
    baja un proveedor lo saca del alta de gastos de todos.
    """
    _assert_general_admin(request)
    used = supabase.table("expenses").select("id").eq("vendor_id", vendor_id).limit(1).execute().data
    if used:
        result = supabase.table("vendors").update(
            {"active": False, "updated_at": _now_iso()}
        ).eq("id", vendor_id).execute()
        if not result.data:
            raise HTTPException(status_code=404, detail="Vendor not found")
        return {"ok": True, "deactivated": True}

    result = supabase.table("vendors").delete().eq("id", vendor_id).execute()
    if not result.data:
        raise HTTPException(status_code=404, detail="Vendor not found")
    return {"ok": True, "deactivated": False}


# ─── Presupuesto ─────────────────────────────────────────────────────────────
_LINE_SELECT = "*, competitions(name), departments(label), accounts(label, kind)"


def _shape_line(row: dict) -> dict:
    comp = row.pop("competitions", None) or {}
    dept = row.pop("departments", None) or {}
    acc = row.pop("accounts", None) or {}
    row["competition_name"] = comp.get("name")
    row["department_label"] = dept.get("label")
    row["account_label"] = acc.get("label")
    return row


@router.get("/lines")
def list_budget_lines(
    request: Request,
    year: int = Query(...),
    department: Optional[str] = Query(None),
    account: Optional[str] = Query(None),
    competition_id: Optional[str] = Query(None),
    general_only: bool = Query(False),
    kind: Optional[str] = Query(None),
):
    q = _scoped(
        supabase.table("budget_lines").select(_LINE_SELECT).eq("year", year).order("account_code"),
        request, department,
    )
    if account:
        q = q.eq("account_code", account)
    if kind:
        q = q.eq("kind", kind)
    if general_only:
        q = q.is_("competition_id")
    elif competition_id:
        q = q.eq("competition_id", competition_id)
    return [_shape_line(r) for r in (q.execute().data or [])]


@router.get("/lines/series/{series_id}")
def line_series(series_id: str, request: Request):
    """La misma línea a lo largo de los años — la vista multi-año de IT."""
    return _scoped(
        supabase.table("budget_lines")
        .select("id, year, amount, escalation_pct, description, department_code")
        .eq("series_id", series_id)
        .order("year"),
        request,
    ).execute().data or []


@router.post("/lines", status_code=201, dependencies=[Depends(require_edit("budget"))])
def create_budget_line(data: BudgetLineCreate, request: Request):
    if not _valid_department(data.department_code):
        raise HTTPException(status_code=400, detail="Unknown or inactive department")
    _assert_can_edit(request, data.department_code)
    kind = data.kind or "expense"
    if kind not in ("expense", "revenue"):
        raise HTTPException(status_code=400, detail="kind must be 'expense' or 'revenue'")
    _check_account(data.account_code, kind)

    record = {k: v for k, v in data.model_dump().items() if v is not None}
    record["kind"] = kind
    record["created_by"] = _user_id(request)
    if data.escalation_pct is None:
        # Default de la cuenta: los códigos de IT ya traen su % del Excel.
        acc = supabase.table("accounts").select("escalation_pct").eq("code", data.account_code).execute().data
        record["escalation_pct"] = (acc[0].get("escalation_pct") if acc else 0) or 0
    return supabase.table("budget_lines").insert(record).execute().data[0]


@router.patch("/lines/{line_id}", dependencies=[Depends(require_edit("budget"))])
def update_budget_line(line_id: str, data: BudgetLineUpdate, request: Request):
    current = _row_or_403("budget_lines", line_id, request, "edit")

    updates = _patch_updates(data, "budget_lines")
    if "department_code" in updates:
        if not _valid_department(updates["department_code"]):
            raise HTTPException(status_code=400, detail="Unknown or inactive department")
        # Mover una línea a otro departamento exige permiso sobre el destino
        # además del origen; si no, se podría empujar gasto a un área ajena.
        _assert_can_edit(request, updates["department_code"])
    kind = updates.get("kind", current.get("kind"))
    if kind not in ("expense", "revenue"):
        raise HTTPException(status_code=400, detail="kind must be 'expense' or 'revenue'")
    if "account_code" in updates:
        _check_account(updates["account_code"], kind)

    # Tocar el monto de una línea calculada la vuelve manual: el recálculo por
    # headcount no puede pisar lo que alguien ajustó a mano.
    if "amount" in updates and current.get("source") == "calculated" and "source" not in updates:
        updates["source"] = "manual"

    updates["updated_at"] = _now_iso()
    result = supabase.table("budget_lines").update(updates).eq("id", line_id).execute()
    if not result.data:
        raise HTTPException(status_code=404, detail="Budget line not found")
    return result.data[0]


@router.delete("/lines/{line_id}", dependencies=[Depends(require_edit("budget"))])
def delete_budget_line(line_id: str, request: Request):
    _row_or_403("budget_lines", line_id, request, "edit")
    result = supabase.table("budget_lines").delete().eq("id", line_id).execute()
    if not result.data:
        raise HTTPException(status_code=404, detail="Budget line not found")
    return {"ok": True}


@router.post("/lines/project", dependencies=[Depends(require_edit("budget"))])
def project_budget(data: BudgetProject, request: Request):
    """Genera los años siguientes aplicando la escalación compuesta.

    Es lo que el Excel de IT hace con las columnas 2028/2029/2030: cada línea
    crece a su % anual sobre el año base. Acá el resultado queda como filas
    reales y editables — proyectar no es un cálculo al vuelo, es un borrador
    que después alguien ajusta.

    Idempotente: una serie ya proyectada se saltea, salvo `overwrite`.
    """
    if not data.to_years:
        raise HTTPException(status_code=400, detail="to_years is required")
    if any(y == data.from_year for y in data.to_years):
        raise HTTPException(status_code=400, detail="to_years cannot include from_year")

    # Proyectar crea filas: sin un departamento explícito haría falta poder
    # editar todos, así que se exige elegir uno salvo acceso total.
    editable = _departments_for(request, "edit")
    if data.department_code:
        _assert_can_edit(request, data.department_code)
    elif editable is not None:
        raise HTTPException(
            status_code=403,
            detail="Pick a department to project — you cannot edit every department",
        )

    q = _scoped(
        supabase.table("budget_lines").select("*").eq("year", data.from_year),
        request, data.department_code,
    )
    source = q.execute().data or []
    if not source:
        raise HTTPException(status_code=404, detail=f"No budget lines for {data.from_year}")

    # Scopeada igual que la de origen: una serie normalmente vive en un solo
    # departamento, pero si alguna vez una de sus filas se movió, `overwrite`
    # pisaría un año que el caller no puede editar.
    existing = _scoped(
        supabase.table("budget_lines")
        .select("id, series_id, year")
        .in_("series_id", [s["series_id"] for s in source]),
        request, data.department_code,
    ).execute().data or []
    by_series_year = {(e["series_id"], e["year"]): e["id"] for e in existing}

    created, updated, skipped = 0, 0, 0
    uid = _user_id(request)

    for target_year in sorted(data.to_years):
        span = target_year - data.from_year
        for line in source:
            pct = float(line.get("escalation_pct") or 0)
            base = float(line.get("amount") or 0)
            # Compuesta sobre el año base, no encadenada año a año: proyectar
            # 2027→2029 directo da lo mismo que 2027→2028→2029, así que el
            # resultado no depende del orden en que se corra.
            amount = round(base * ((1 + pct) ** span), 2)

            key = (line["series_id"], target_year)
            if key in by_series_year:
                if not data.overwrite:
                    skipped += 1
                    continue
                supabase.table("budget_lines").update(
                    {"amount": amount, "updated_at": _now_iso()}
                ).eq("id", by_series_year[key]).execute()
                updated += 1
                continue

            record = {
                "year": target_year,
                "department_code": line["department_code"],
                "account_code": line["account_code"],
                "competition_id": line.get("competition_id"),
                "kind": line.get("kind", "expense"),
                "description": line["description"],
                "qty": line.get("qty"),
                "monthly_amount": round(amount / 12, 2) if line.get("monthly_amount") else None,
                "amount": amount,
                "escalation_pct": pct,
                "source": line.get("source", "manual"),
                "notes": line.get("notes"),
                "series_id": line["series_id"],
                "created_by": uid,
            }
            record = {k: v for k, v in record.items() if v is not None}
            supabase.table("budget_lines").insert(record).execute()
            created += 1

    return {"created": created, "updated": updated, "skipped": skipped}


# ─── Import / export del presupuesto en Excel ────────────────────────────────
#
# Dos pasos como el resto de los importadores del repo: el preview no escribe
# nada y muestra qué línea se crea, cuál se actualiza y qué columna quedó sin
# resolver; recién /commit toca la base. Lo dudoso NUNCA se importa solo —
# imputar una columna al evento equivocado es plata en el presupuesto de otro.
#
# Autorización: `require_edit("budget")` como cualquier escritura, más el
# recorte por departamento. Una planilla puede tocar varios departamentos (las
# líneas de Comms e IT viven dentro del presupuesto de Competitions), así que
# se valida el elegido acá y el servicio marca como error las filas de un
# departamento que el caller no puede editar.

_MAX_SHEET_BYTES = 5 * 1024 * 1024


async def _read_sheet_upload(file: UploadFile) -> bytes:
    if not (file.filename or "").lower().endswith((".xlsx", ".xlsm")):
        raise HTTPException(status_code=400, detail="Only .xlsx files are supported")
    content = await file.read()
    if len(content) > _MAX_SHEET_BYTES:
        raise HTTPException(status_code=400, detail="File too large (max 5 MB)")
    return content


def _import_options(request: Request, year: int, department: str, kind: str,
                    sheet, mapping, default_competition_id,
                    create_accounts: bool, replace: bool) -> dict:
    if kind not in ("expense", "revenue"):
        raise HTTPException(status_code=400, detail="kind must be 'expense' or 'revenue'")
    if not 2000 <= year <= 2100:
        raise HTTPException(status_code=400, detail="Invalid year")
    if not _valid_department(department):
        raise HTTPException(status_code=400, detail="Unknown or inactive department")
    _assert_can_edit(request, department)

    if default_competition_id:
        try:
            uuid.UUID(default_competition_id)
        except ValueError:
            raise HTTPException(status_code=400, detail="Invalid competition id")
        if not supabase.table("competitions").select("id").eq("id", default_competition_id).execute().data:
            raise HTTPException(status_code=400, detail="Unknown competition")

    parsed_mapping = {}
    if mapping:
        try:
            parsed_mapping = json.loads(mapping)
        except ValueError:
            raise HTTPException(status_code=400, detail="Invalid column mapping")
        if not isinstance(parsed_mapping, dict) or not all(
            isinstance(k, str) and isinstance(v, str) for k, v in parsed_mapping.items()
        ):
            raise HTTPException(status_code=400, detail="Invalid column mapping")

    return {
        "year": year,
        "kind": kind,
        "department": department,
        "sheet": sheet or None,
        "mapping": parsed_mapping,
        "default_competition_id": default_competition_id or None,
        "create_accounts": create_accounts,
        "replace": replace,
        "editable": _departments_for(request, "edit"),
        "viewable": _departments_for(request, "view"),
    }


@router.post("/import/lines/preview", dependencies=[Depends(require_edit("budget"))])
async def preview_line_import(
    request: Request,
    file: UploadFile = File(...),
    year: int = Form(...),
    department: str = Form(...),
    kind: str = Form("expense"),
    sheet: Optional[str] = Form(None),
    mapping: Optional[str] = Form(None),
    default_competition_id: Optional[str] = Form(None),
    create_accounts: bool = Form(True),
    replace: bool = Form(False),
):
    from api._lib.services import budget_import

    content = await _read_sheet_upload(file)
    options = _import_options(request, year, department, kind, sheet, mapping,
                              default_competition_id, create_accounts, replace)
    try:
        return budget_import.preview(content, **options)
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc))


@router.post("/import/lines/commit", dependencies=[Depends(require_edit("budget"))])
async def commit_line_import(
    request: Request,
    file: UploadFile = File(...),
    year: int = Form(...),
    department: str = Form(...),
    kind: str = Form("expense"),
    sheet: Optional[str] = Form(None),
    mapping: Optional[str] = Form(None),
    default_competition_id: Optional[str] = Form(None),
    create_accounts: bool = Form(True),
    replace: bool = Form(False),
):
    from api._lib.services import budget_import

    content = await _read_sheet_upload(file)
    options = _import_options(request, year, department, kind, sheet, mapping,
                              default_competition_id, create_accounts, replace)
    try:
        result = budget_import.commit(content, created_by=_user_id(request), **options)
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc))
    logger.info(
        "budget import: %s created, %s updated, %s deleted (year=%s, departments=%s, user=%s)",
        result["created"], result["updated"], result["deleted"],
        year, ",".join(result["departments"]), _user_id(request),
    )
    return result


@router.get("/lines/export.xlsx")
def export_budget_lines(
    request: Request,
    year: int = Query(...),
    department: Optional[str] = Query(None),
    kind: str = Query("expense"),
    lang: str = Query("es"),
):
    """Las líneas de un año en el mismo formato que lee el import.

    Es la plantilla y el round-trip a la vez: se baja, se edita en Excel y se
    vuelve a subir. La columna ID hace que la reimportación actualice esas
    mismas filas en vez de duplicarlas.
    """
    from api._lib.services.budget_import import export_lines_xlsx

    query = _scoped(
        supabase.table("budget_lines").select(_LINE_SELECT).eq("year", year)
        .eq("kind", kind).order("account_code"),
        request, department,
    )
    lines = [_shape_line(r) for r in (query.execute().data or [])]
    content = export_lines_xlsx(lines, year, "en" if lang == "en" else "es")
    return Response(
        content=content,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="budget-{year}.xlsx"'},
    )


# `split` dejó de existir en la migración 047: el pago parcial es un
# `in_process` con `paid_amount` (ver _payment_buckets).
_OPEN_PAYMENT_STATUSES = ("new", "in_process")


def _year_of(value) -> Optional[int]:
    """Año de una fecha ISO ('2026-03-01' / timestamptz), o None."""
    text = str(value or "")
    return int(text[:4]) if len(text) >= 4 and text[:4].isdigit() else None


def _payment_fiscal_year(p: dict) -> Optional[int]:
    """A qué año del presupuesto se imputa un pago. None = a ninguno.

    - `completed` → el año de `payment_date`, que es cuando salió la plata. Si
      no tiene fecha (la UI de Payments nunca la pidió), cae al año de la
      competencia: antes un completado sin fecha desaparecía de TODOS los años.
    - abierto (`new` / `in_process`) → el año de la competencia, que
      es contra qué presupuesto compromete. Antes se contaba como comprometido
      en cualquier año que se consultara.
    - cualquier otro estado → no cuenta.

    El año de la competencia es `competitions.year`, y si falta, el de
    `start_date`. Espera la fila con el embed `nominations(competitions(...))`.
    """
    comp = ((p.get("nominations") or {}).get("competitions")) or {}
    comp_year = comp.get("year") or _year_of(comp.get("start_date"))
    status = p.get("status")
    if status == "completed":
        return _year_of(p.get("payment_date")) or comp_year
    if status in _OPEN_PAYMENT_STATUSES:
        return comp_year
    return None


def _payment_buckets(status: Optional[str], amount, paid_amount) -> tuple:
    """(ejecutado, comprometido) de un monto de pago según su estado.

    - `completed`  → todo ejecutado.
    - `in_process` con `paid_amount > 0` → lo pagado ejecutado (con tope en el
      monto) y el resto comprometido: es el pago parcial que antes escondía
      `split`.
    - `new` / `in_process` sin pagar → todo comprometido.
    - `cancelled` u otro → nada.

    Para el pasaje se llama con `paid_amount=None`: lo pagado es del fee, y el
    pasaje recién cuenta como ejecutado al completarse.
    """
    total = float(amount or 0)
    if status == "completed":
        return round(total, 2), 0.0
    if status not in _OPEN_PAYMENT_STATUSES:
        return 0.0, 0.0
    paid = float(paid_amount or 0) if status == "in_process" else 0.0
    paid = min(max(paid, 0.0), total)
    return round(paid, 2), round(total - paid, 2)


def _revenue_year(r: dict) -> Optional[int]:
    """A qué año se imputa un ingreso. None = a ninguno.

    Recibido → `received_date` (el CHECK revenue_received_has_date garantiza
    que exista); esperado → `expected_date`; cancelado o sin fecha → afuera.
    Un esperado sin fecha no tiene año al que pertenecer: la UI lo marca para
    que se complete en vez de sumarlo a cualquier año que se consulte.
    """
    status = r.get("status")
    if status == "received":
        return _year_of(r.get("received_date"))
    if status == "expected":
        return _year_of(r.get("expected_date"))
    return None


# Orden fijo de los desgloses por estado: el front y el Excel los muestran en
# este orden aunque un estado no tenga filas.
_EXPENSE_STATUS_ORDER = ("draft", "approved", "paid", "rejected", "cancelled")
_PAYMENT_STATUS_ORDER = ("new", "in_process", "completed", "cancelled")


def _month_of(value) -> Optional[str]:
    """'YYYY-MM' de una fecha ISO ('2026-03-14' / timestamptz), o None."""
    text = str(value or "")[:7]
    return text if _PERIOD_RE.match(text) else None


def _payment_status_year(p: dict) -> Optional[int]:
    """Año en el que un pago aparece en `by_status` del summary.

    Para los estados que imputan es el mismo de _payment_fiscal_year. Un
    `cancelled` no imputa a ningún año, pero en el desglose por estado sí
    tiene que contarse: va al año de su `payment_date` y, si no tiene, al de la
    competencia. Si tampoco hay competencia con año, no aparece en ninguno.
    """
    fiscal = _payment_fiscal_year(p)
    if fiscal is not None:
        return fiscal
    comp = ((p.get("nominations") or {}).get("competitions")) or {}
    return _year_of(p.get("payment_date")) or comp.get("year") or _year_of(comp.get("start_date"))


def _expense_month(e: dict, year: int) -> Optional[str]:
    """Mes del año pedido al que va un gasto en `by_month`.

    - `paid` → el mes de `payment_date`, que es cuando salió la plata. Si la
      fecha de pago cae en otro año (factura de diciembre pagada en enero), se
      usa la del gasto: el gasto imputa al año de `expense_date` y, si se lo
      mandara a un mes de otro año, la suma de los 12 meses dejaría de cerrar
      con el ejecutado del año.
    - `approved` → el mes de `expense_date`.
    """
    if e.get("status") == "paid":
        paid = _month_of(e.get("payment_date"))
        if paid and paid.startswith(f"{year}-"):
            return paid
    return _month_of(e.get("expense_date"))


@router.get("/summary")
def budget_summary(request: Request, year: int = Query(...), department: Optional[str] = Query(None)):
    """Presupuestado vs. comprometido vs. ejecutado vs. restante (ver _summary_data)."""
    return _summary_data(request, year, department)


def _summary_data(request: Request, year: int, department: Optional[str] = None) -> dict:
    """Presupuestado vs. comprometido vs. ejecutado vs. restante.

    Es el cuerpo de GET /summary, separado para que GET /summary/export.xlsx
    exporte exactamente lo mismo que se ve en pantalla. Se llama con valores
    de verdad, nunca con los defaults `Query()` (ver _fetch_expenses).

    Desde la migración 036 el ejecutado suma **las dos** fuentes de gasto:
    `expenses` y los pagos a personas nominadas (`payments`, que ya llevan
    departamento y cuenta). Un pago cuenta como ejecutado cuando su status es
    `completed` — es el equivalente de `paid` en el enum viejo de payments.

    Los pagos sin imputar (sin departamento) se reportan aparte en
    `unallocated_payments` en vez de repartirse a ciegas: así se ve qué falta
    clasificar en lugar de esconderlo dentro de un total.

    Un pago aporta DOS imputaciones, no una: el fee (`total`) a la cuenta de
    fees del rol y el pasaje (`airfare`) a la de travel. Son líneas distintas
    del presupuesto — COMP-11 vs. COMP-12 para un TD — y meter el pasaje en la
    del fee infla una y deja la otra en cero para siempre.
    """
    lq = _scoped(supabase.table("budget_lines").select("*").eq("year", year), request, department)
    # Gastos del año: SOLO el .gte va a la base. El query builder guarda un
    # filtro por columna (database.py), así que encadenar .gte().lte() sobre
    # expense_date dejaba vivo únicamente el .lte: cada año sumaba también
    # todos los anteriores (los $41.667 de IT 2026 aparecían en 2027). El tope
    # se aplica en Python, abajo — mismo arreglo que _fetch_expenses.
    eq_ = _scoped(
        supabase.table("expenses")
        .select("amount, status, department_code, account_code, competition_id, expense_date, payment_date")
        .gte("expense_date", f"{year}-01-01"),
        request, department,
    )
    # Ingresos: el año sale de dos columnas distintas según el estado, y
    # PostgREST no hace OR entre columnas con este builder; se filtra en
    # Python (ver _revenue_year). Antes no se filtraba nada y el dashboard de
    # cada año sumaba los ingresos de todos los años.
    rq = _scoped(
        supabase.table("revenues").select(
            "amount, status, department_code, competition_id, expected_date, received_date"
        ),
        request, department,
    )
    pq = _scoped(
        supabase.table("payments").select(
            "total, airfare, paid_amount, status, department_code, account_code, payment_date, "
            "nominations(competition_id, personnel(role), competitions(year, start_date))"
        ),
        request, department,
    )

    lines = lq.execute().data or []
    year_end = f"{year}-12-31"
    expenses = [e for e in (eq_.execute().data or [])
                if str(e.get("expense_date") or "")[:10] <= year_end]
    revenues = [r for r in (rq.execute().data or []) if _revenue_year(r) == year]

    # Los pagos no tienen fecha de gasto: el año sale de _payment_fiscal_year
    # (fecha de pago si está completado, año de la competencia si está
    # abierto). Se filtra en Python porque hace falta mirar el embed.
    #
    # `by_status` cuenta además los que no imputan (cancelados): es el desglose
    # de "qué hay cargado", no de "qué consume presupuesto".
    payment_status = {s: {"count": 0, "amount": 0.0} for s in _PAYMENT_STATUS_ORDER}
    payments = []
    for p in (pq.execute().data or []):
        if _payment_status_year(p) == year:
            entry = payment_status.setdefault(p.get("status") or "—", {"count": 0, "amount": 0.0})
            entry["count"] += 1
            # Fee + pasaje: lo mismo que el pago aporta a los totales.
            entry["amount"] += float(p.get("total") or 0) + float(p.get("airfare") or 0)
        if _payment_fiscal_year(p) != year:
            continue
        nom = p.pop("nominations", None) or {}
        # Mes de cada fila para `by_month`. Lo comprometido va al mes en que
        # arranca la competencia (es cuándo se va a pagar); lo ejecutado, al de
        # `payment_date`, y si no la tiene o cae en otro año, también al de la
        # competencia. Sin `start_date` el mes queda en None: la fila suma a los
        # totales pero no a ningún mes.
        start_month = _month_of((nom.get("competitions") or {}).get("start_date"))
        paid_month = _month_of(p.get("payment_date"))
        if not (paid_month and paid_month.startswith(f"{year}-")):
            paid_month = start_month
        base = {
            "status": p.get("status"),
            "department_code": p.get("department_code"),
            "competition_id": nom.get("competition_id"),
        }
        # Cada fila lleva su `bucket` (executed / committed). Un fee
        # `in_process` con `paid_amount` (pago parcial, migración 047) se parte
        # en dos filas: lo pagado es ejecutado y el resto comprometido.
        fee_exec, fee_open = _payment_buckets(p.get("status"), p.get("total"), p.get("paid_amount"))
        for bucket, amount in (("executed", fee_exec), ("committed", fee_open)):
            if amount:
                payments.append({**base, "bucket": bucket, "amount": amount,
                                 "account_code": p.get("account_code"),
                                 "month": paid_month if bucket == "executed" else start_month})
        # El pasaje se imputa a la cuenta de travel del rol, no a la del fee.
        # Se agrega como una fila más para que todo lo de abajo —ejecutado,
        # comprometido, los tres rollups, el sin-imputar— lo cuente sin
        # enterarse de que existe un segundo concepto. `paid_amount` es del
        # fee: el pasaje sigue entero como comprometido hasta `completed`.
        if float(p.get("airfare") or 0):
            role = (nom.get("personnel") or {}).get("role")
            air_exec, air_open = _payment_buckets(p.get("status"), p.get("airfare"), None)
            payments.append({**base, "bucket": "executed" if air_exec else "committed",
                             "amount": air_exec or air_open,
                             "account_code": travel_account_for_role(role),
                             "month": paid_month if air_exec else start_month})

    depts = {d["code"]: d["label"] for d in (supabase.table("departments").select("code, label").execute().data or [])}
    accts = {a["code"]: a["label"] for a in (supabase.table("accounts").select("code, label").execute().data or [])}
    comps = {c["id"]: c["name"] for c in (supabase.table("competitions").select("id, name").execute().data or [])}

    def _num(row, field="amount") -> float:
        return float(row.get(field) or 0)

    exp_lines = [l for l in lines if l.get("kind") == "expense"]
    rev_lines = [l for l in lines if l.get("kind") == "revenue"]

    # Un pago `completed` es ejecutado; uno abierto es comprometido, salvo lo
    # ya pagado de un `in_process` parcial (ver _payment_buckets). `cancelled`
    # no llega hasta acá (_payment_fiscal_year lo deja afuera).
    paid_payments = [p for p in payments if p.get("bucket") == "executed"]
    open_payments = [p for p in payments if p.get("bucket") == "committed"]

    budgeted = round(sum(_num(l) for l in exp_lines), 2)
    executed = round(
        sum(_num(e) for e in expenses if e.get("status") == "paid")
        + sum(_num(p) for p in paid_payments), 2)
    committed = round(
        sum(_num(e) for e in expenses if e.get("status") == "approved")
        + sum(_num(p) for p in open_payments), 2)
    unallocated = round(sum(_num(p) for p in payments if not p.get("department_code")), 2)

    def _rollup(key: str, labels: dict, is_competition: bool = False) -> list:
        acc: dict = {}

        def _entry(k):
            if k is None:
                name = None if not is_competition else "—"
                return acc.setdefault("__general__", {
                    "key": None, "label": name, "general": True,
                    "budgeted": 0.0, "executed": 0.0, "committed": 0.0,
                })
            return acc.setdefault(k, {
                "key": k, "label": labels.get(k, k), "general": False,
                "budgeted": 0.0, "executed": 0.0, "committed": 0.0,
            })

        for line in exp_lines:
            _entry(line.get(key))["budgeted"] += _num(line)
        for e in expenses:
            entry = _entry(e.get(key))
            if e.get("status") == "paid":
                entry["executed"] += _num(e)
            elif e.get("status") == "approved":
                entry["committed"] += _num(e)
        for p in paid_payments:
            _entry(p.get(key))["executed"] += _num(p)
        for p in open_payments:
            _entry(p.get(key))["committed"] += _num(p)

        out = []
        for entry in acc.values():
            for f in ("budgeted", "executed", "committed"):
                entry[f] = round(entry[f], 2)
            entry["remaining"] = round(entry["budgeted"] - entry["executed"], 2)
            out.append(entry)
        return sorted(out, key=lambda e: -e["budgeted"])

    # Serie mensual del año: siempre 12 meses, en cero los que no tienen
    # movimiento, para que el gráfico acumulado no tenga huecos. Sale de las
    # mismas filas que los totales; lo único que puede quedar afuera es un pago
    # sin fecha útil (ver el loop de pagos).
    months = {f"{year}-{m:02d}": {"executed": 0.0, "committed": 0.0} for m in range(1, 13)}
    for e in expenses:
        bucket = {"paid": "executed", "approved": "committed"}.get(e.get("status"))
        month = _expense_month(e, year) if bucket else None
        if month in months:
            months[month][bucket] += _num(e)
    for p in payments:
        if p.get("month") in months:
            months[p["month"]][p["bucket"]] += _num(p)
    by_month = [
        {"month": m, "executed": round(v["executed"], 2), "committed": round(v["committed"], 2)}
        for m, v in months.items()
    ]

    # Desglose por estado sobre las filas del año, incluidos borradores,
    # rechazados y cancelados, que no entran a ningún total.
    expense_status = {s: {"count": 0, "amount": 0.0} for s in _EXPENSE_STATUS_ORDER}
    for e in expenses:
        entry = expense_status.setdefault(e.get("status") or "—", {"count": 0, "amount": 0.0})
        entry["count"] += 1
        entry["amount"] += _num(e)
    for group in (expense_status, payment_status):
        for entry in group.values():
            entry["amount"] = round(entry["amount"], 2)

    return {
        "year": year,
        # Desde la 036 los pagos a personas SÍ se cuentan. El flag se mantiene
        # para que un frontend viejo no mienta en pantalla.
        "excludes_person_payments": False,
        "unallocated_payments": unallocated,
        "totals": {
            "budgeted": budgeted,
            "executed": executed,
            "committed": committed,
            # Restante contra lo ejecutado: aprobado-sin-pagar NO descuenta.
            "remaining": round(budgeted - executed, 2),
            "revenue_budgeted": round(sum(_num(l) for l in rev_lines), 2),
            "revenue_received": round(sum(_num(r) for r in revenues if r.get("status") == "received"), 2),
            "revenue_expected": round(sum(_num(r) for r in revenues if r.get("status") == "expected"), 2),
        },
        "by_department": _rollup("department_code", depts),
        "by_account": _rollup("account_code", accts),
        "by_competition": _rollup("competition_id", comps, is_competition=True),
        "by_month": by_month,
        "by_status": {"expenses": expense_status, "payments": payment_status},
    }


_SUMMARY_EXPORT_TEXT = {
    "es": {
        "sheets": ("Totales", "Por departamento", "Por cuenta", "Por competencia", "Por mes"),
        "concept": "Concepto", "amount": "Monto",
        "totals": (
            ("budgeted", "Presupuestado"), ("executed", "Ejecutado"),
            ("committed", "Comprometido"), ("remaining", "Restante"),
            ("revenue_budgeted", "Ingresos presupuestados"),
            ("revenue_received", "Ingresos recibidos"),
            ("revenue_expected", "Ingresos esperados"),
        ),
        "unallocated": "Pagos sin imputar",
        "code": "Código", "department": "Departamento", "account": "Cuenta",
        "competition": "Competencia", "month": "Mes",
        "budgeted": "Presupuestado", "executed": "Ejecutado", "committed": "Comprometido",
        "remaining": "Restante", "pct": "% ejecutado", "cumulative": "Ejecutado acumulado",
        "no_key": "Sin imputar", "general": "General", "total": "Total",
    },
    "en": {
        "sheets": ("Totals", "By department", "By account", "By competition", "By month"),
        "concept": "Item", "amount": "Amount",
        "totals": (
            ("budgeted", "Budgeted"), ("executed", "Executed"),
            ("committed", "Committed"), ("remaining", "Remaining"),
            ("revenue_budgeted", "Budgeted revenue"),
            ("revenue_received", "Revenue received"),
            ("revenue_expected", "Revenue expected"),
        ),
        "unallocated": "Unallocated payments",
        "code": "Code", "department": "Department", "account": "Account",
        "competition": "Competition", "month": "Month",
        "budgeted": "Budgeted", "executed": "Executed", "committed": "Committed",
        "remaining": "Remaining", "pct": "% executed", "cumulative": "Cumulative executed",
        "no_key": "Unallocated", "general": "General", "total": "Total",
    },
}
_ROLLUP_FIELDS = ("budgeted", "executed", "committed", "remaining")


def _pct(executed: float, budgeted: float) -> Optional[float]:
    return round(executed / budgeted, 4) if budgeted else None


def _write_rollup_sheet(ws, rows: list, t: dict, label_header: str, with_code: bool,
                        empty_label: str) -> None:
    """Una hoja de rollup (departamento / cuenta / competencia) con fila de total.

    Las filas vienen de `_rollup()` en _summary_data, ya ordenadas por
    presupuestado. Sin código para competencias: la key es un UUID.
    """
    from api._lib.services.xlsx_utils import money_format, safe_text, style_header

    lead = [t["code"], label_header] if with_code else [label_header]
    ws.append(lead + [t[f] for f in _ROLLUP_FIELDS] + [t["pct"]])
    style_header(ws)
    sums = dict.fromkeys(_ROLLUP_FIELDS, 0.0)
    for r in rows:
        label = r.get("label") or (empty_label if r.get("key") is None else r.get("key"))
        if r.get("general") and not with_code:
            label = t["general"]
        values = [float(r.get(f) or 0) for f in _ROLLUP_FIELDS]
        for f, v in zip(_ROLLUP_FIELDS, values):
            sums[f] += v
        lead_vals = [safe_text(r.get("key")), safe_text(label)] if with_code else [safe_text(label)]
        ws.append(lead_vals + values + [_pct(values[1], values[0])])
    total_lead = [t["total"], None] if with_code else [t["total"]]
    ws.append(total_lead + [round(sums[f], 2) for f in _ROLLUP_FIELDS]
              + [_pct(sums["executed"], sums["budgeted"])])
    _bold_last_row(ws)

    first_money = len(lead) + 1
    money_format(ws, range(first_money, first_money + len(_ROLLUP_FIELDS)))
    money_format(ws, [first_money + len(_ROLLUP_FIELDS)], fmt="0.0%")


@router.get("/summary/export.xlsx")
def export_summary(
    request: Request,
    year: int = Query(...),
    department: Optional[str] = Query(None),
    lang: str = Query("es"),
):
    """El dashboard de un año en Excel: Totales, Por departamento, Por cuenta,
    Por competencia y Por mes. Sale de _summary_data, así que el recorte por
    departamento (_scoped) es el mismo que en pantalla."""
    from openpyxl import Workbook
    from api._lib.services.xlsx_utils import (
        autosize, money_format, style_header, workbook_bytes, xlsx_response,
    )

    data = _summary_data(request, year, department)
    t = _SUMMARY_EXPORT_TEXT["en" if lang == "en" else "es"]
    s_totals, s_dept, s_acct, s_comp, s_month = t["sheets"]

    wb = Workbook()

    ws = wb.active
    ws.title = s_totals
    ws.append([t["concept"], t["amount"]])
    style_header(ws)
    totals = data.get("totals") or {}
    for key, label in t["totals"]:
        ws.append([label, float(totals.get(key) or 0)])
    ws.append([t["unallocated"], float(data.get("unallocated_payments") or 0)])
    money_format(ws, [2])
    autosize(ws)

    for title, key, header, with_code, empty in (
        (s_dept, "by_department", t["department"], True, t["no_key"]),
        (s_acct, "by_account", t["account"], True, t["no_key"]),
        (s_comp, "by_competition", t["competition"], False, t["general"]),
    ):
        ws = wb.create_sheet(title)
        _write_rollup_sheet(ws, data.get(key) or [], t, header, with_code, empty)
        autosize(ws)

    ws = wb.create_sheet(s_month)
    ws.append([t["month"], t["executed"], t["committed"], t["cumulative"]])
    style_header(ws)
    cumulative = sum_exec = sum_comm = 0.0
    for m in data.get("by_month") or []:
        executed, committed = float(m.get("executed") or 0), float(m.get("committed") or 0)
        cumulative += executed
        sum_exec += executed
        sum_comm += committed
        ws.append([m.get("month"), executed, committed, round(cumulative, 2)])
    ws.append([t["total"], round(sum_exec, 2), round(sum_comm, 2), None])
    _bold_last_row(ws)
    money_format(ws, [2, 3, 4])
    autosize(ws)

    suffix = f"-{department}" if department else ""
    return xlsx_response(workbook_bytes(wb), f"budget-summary-{year}{suffix}.xlsx")


@router.get("/competitions/{competition_id}/cost")
def competition_cost(competition_id: str, request: Request):
    """Costo total de una competencia: pagos a personas + airfare + gastos.

    ⚠️ EXCEPCIÓN EXPLÍCITA al recorte por departamento — este endpoint NO llama
    a _scoped(). Es deliberado y fue decisión del cliente: el scoping rige
    cargar y editar, no leer el costo de un evento. Una competencia cruza
    departamentos (la línea de TV Production es de Comms dentro del presupuesto
    de Competitions), así que sin el desglose entero la cifra no sirve para
    reportar. Quien tiene el módulo `budget` ve el evento completo.

    Junta las tres fuentes que hoy están separadas:
      - `payments`  → fees de las personas nominadas, + airfare aparte
      - `expenses`  → gasto del evento sin persona (shipping, branding, seguros)
      - `budget_lines` → contra qué se compara

    El airfare va como línea propia y NO dentro del fee: en `payments` el total
    es amount + extra y el vuelo se liquida con la agencia (migración 013). Pero
    línea propia no es plata invisible: cuenta como ejecutado igual que el fee y
    entra en el desglose por departamento — lo que cambia es la cuenta (la de
    travel del rol, ver `budget_accounts.py`).
    """
    comp = supabase.table("competitions").select("id, name, year, subzone, tier") \
        .eq("id", competition_id).execute().data
    if not comp:
        raise HTTPException(status_code=404, detail="Competition not found")
    comp = comp[0]

    lines = supabase.table("budget_lines").select(
        "amount, department_code, account_code, kind"
    ).eq("competition_id", competition_id).execute().data or []

    expenses = supabase.table("expenses").select(
        "amount, status, department_code, account_code, description"
    ).eq("competition_id", competition_id).execute().data or []

    noms = supabase.table("nominations").select("id").eq("competition_id", competition_id) \
        .execute().data or []
    pays = []
    if noms:
        pays = supabase.table("payments").select(
            "total, airfare, paid_amount, status, department_code, account_code, "
            "nominations(personnel(name, role))"
        ).in_("nomination_id", [n["id"] for n in noms]).execute().data or []

    def _n(row, field="amount") -> float:
        return float(row.get(field) or 0)

    budgeted = round(sum(_n(l) for l in lines if l.get("kind") == "expense"), 2)
    exp_paid = round(sum(_n(e) for e in expenses if e.get("status") == "paid"), 2)
    exp_open = round(sum(_n(e) for e in expenses if e.get("status") == "approved"), 2)
    # Fee: un `in_process` parcial aporta lo pagado a ejecutado y el resto a
    # comprometido; pasaje: entero comprometido hasta `completed`. Mismo
    # criterio que /summary (_payment_buckets). `cancelled` no cuenta.
    fee_split = [_payment_buckets(p.get("status"), p.get("total"), p.get("paid_amount")) for p in pays]
    air_split = [_payment_buckets(p.get("status"), p.get("airfare"), None) for p in pays]
    fees_paid = round(sum(e for e, _ in fee_split), 2)
    fees_open = round(sum(c for _, c in fee_split), 2)
    airfare_paid = round(sum(e for e, _ in air_split), 2)
    airfare_open = round(sum(c for _, c in air_split), 2)
    airfare = round(airfare_paid + airfare_open, 2)

    executed = round(exp_paid + fees_paid + airfare_paid, 2)

    # Desglose por departamento, con las tres fuentes juntas.
    by_dept: dict = {}

    def _entry(code):
        return by_dept.setdefault(code or "—", {
            "department_code": code, "budgeted": 0.0, "executed": 0.0, "committed": 0.0,
        })

    for l in lines:
        if l.get("kind") == "expense":
            _entry(l.get("department_code"))["budgeted"] += _n(l)
    for e in expenses:
        entry = _entry(e.get("department_code"))
        if e.get("status") == "paid":
            entry["executed"] += _n(e)
        elif e.get("status") == "approved":
            entry["committed"] += _n(e)
    for p in pays:
        entry = _entry(p.get("department_code"))
        # Fee y pasaje van al mismo departamento (cambia la cuenta, y acá el
        # desglose es por departamento), así que se suman juntos.
        fee_exec, fee_open = _payment_buckets(p.get("status"), p.get("total"), p.get("paid_amount"))
        air_exec, air_open = _payment_buckets(p.get("status"), p.get("airfare"), None)
        entry["executed"] += fee_exec + air_exec
        entry["committed"] += fee_open + air_open

    labels = {d["code"]: d["label"] for d in
              (supabase.table("departments").select("code, label").execute().data or [])}
    dept_rows = []
    for code, e in by_dept.items():
        for f in ("budgeted", "executed", "committed"):
            e[f] = round(e[f], 2)
        e["label"] = labels.get(e["department_code"], e["department_code"] or "—")
        e["remaining"] = round(e["budgeted"] - e["executed"], 2)
        dept_rows.append(e)
    dept_rows.sort(key=lambda e: -max(e["budgeted"], e["executed"]))

    people = []
    for p in pays:
        person = ((p.pop("nominations", None) or {}).get("personnel")) or {}
        people.append({
            "name": person.get("name"), "role": person.get("role"),
            "total": p.get("total"), "airfare": p.get("airfare"), "status": p.get("status"),
        })
    people.sort(key=lambda r: (r["name"] or "").lower())

    return {
        "competition": comp,
        "scoped": False,   # la UI lo dice: acá se ve el evento completo
        "totals": {
            "budgeted": budgeted,
            "executed": executed,
            "committed": round(exp_open + fees_open + airfare_open, 2),
            "remaining": round(budgeted - executed, 2),
            "person_fees": round(fees_paid + fees_open, 2),
            "airfare": airfare,
            "event_expenses": round(exp_paid + exp_open, 2),
            "people_count": len(people),
        },
        "by_department": dept_rows,
        "people": people,
    }


# ─── Headcount y supuestos (base de las líneas calculadas — UI en fase 5) ────
@router.get("/headcount")
def list_headcount(year: int = Query(...)):
    return (
        supabase.table("budget_headcount")
        .select("*, competitions(name)")
        .eq("year", year)
        .order("role_label")
        .execute()
        .data
    ) or []


@router.put("/headcount", dependencies=[Depends(require_edit("budget"))])
def put_headcount(data: HeadcountPut, request: Request):
    """Reemplaza el headcount del año. Un headcount en 0 borra la fila.

    El headcount (personas × evento) es el insumo de las líneas de travel de
    Competitions: exige poder editar ese departamento, igual que
    `/headcount/generate`. Con solo `require_edit("budget")` el designado de IT
    podía mover el presupuesto de viaje de otra área.
    """
    _assert_can_edit(request, "competitions")
    for item in data.items:
        existing = supabase.table("budget_headcount").select("id").eq("year", data.year) \
            .eq("role_label", item.role_label)
        existing = (existing.is_("competition_id") if not item.competition_id
                    else existing.eq("competition_id", item.competition_id)).execute().data

        if item.headcount <= 0:
            if existing:
                supabase.table("budget_headcount").delete().eq("id", existing[0]["id"]).execute()
            continue
        payload = {"headcount": item.headcount, "account_code": item.account_code}
        if existing:
            supabase.table("budget_headcount").update(payload).eq("id", existing[0]["id"]).execute()
        else:
            supabase.table("budget_headcount").insert({
                "year": data.year,
                "competition_id": item.competition_id,
                "role_label": item.role_label,
                **payload,
            }).execute()
    return {"ok": True}


@router.post("/headcount/generate", dependencies=[Depends(require_edit("budget"))])
def generate_travel_lines(request: Request, year: int = Query(...),
                          department: str = Query("competitions"),
                          replace_manual: bool = Query(False)):
    """Convierte el headcount en líneas de presupuesto de viaje.

    Replica lo que el Excel hace entre los sheets `Staffing (travel)` y
    `Breakdown`: headcount × costo de vuelo promedio → la línea de travel de
    ese rol. En 2027 son 235 persona-evento × $1.050.

    Las líneas nacen con `source = 'calculated'` y se REGENERAN en cada corrida.
    Las que alguien editó a mano ya pasaron a 'manual' (lo hace
    update_budget_line) y se dejan intactas: el recálculo nunca pisa un ajuste
    humano.

    ⚠️ RIESGO DE DOBLE CONTEO, y por eso el chequeo de conflictos. Las líneas de
    travel del Excel de Competitions (COMP-12 Technical Delegate Travel, COMP-14
    TV Graphics Operator Travel, COMP-19 FA Staff Travel…) YA SON headcount ×
    costo de vuelo: así las calculó quien armó la planilla. Generar encima de un
    año que ya tiene esas líneas cargadas a mano duplicaría el presupuesto de
    viaje. Por defecto esos pares (cuenta, competencia) se saltean y se
    reportan; `replace_manual` los borra y los reemplaza por la versión
    calculada.
    """
    _assert_can_edit(request, department)

    rows = supabase.table("budget_headcount").select("*").eq("year", year).execute().data or []
    rows = [r for r in rows if r.get("account_code") and r.get("headcount")]
    if not rows:
        raise HTTPException(status_code=404, detail=f"No headcount with an account for {year}")

    a = supabase.table("budget_assumptions").select("avg_flight_cost").eq("year", year).execute().data
    flight = float(a[0]["avg_flight_cost"]) if a else 1050.0

    # Solo se borran las calculadas: las manuales sobreviven al regenerado.
    old = supabase.table("budget_lines").select("id, source") \
        .eq("year", year).eq("department_code", department) \
        .eq("source", "calculated").execute().data or []
    for o in old:
        supabase.table("budget_lines").delete().eq("id", o["id"]).execute()

    # Pares (cuenta, competencia) que ya tienen una línea manual: son las que
    # duplicarían el presupuesto de viaje.
    manual = supabase.table("budget_lines").select("id, account_code, competition_id") \
        .eq("year", year).eq("department_code", department) \
        .eq("source", "manual").execute().data or []
    taken = {(m["account_code"], m["competition_id"]): m["id"] for m in manual}

    made, conflicts = [], []
    for r in rows:
        key = (r["account_code"], r.get("competition_id"))
        if key in taken:
            if not replace_manual:
                conflicts.append({"account_code": key[0], "competition_id": key[1],
                                  "role_label": r["role_label"], "headcount": r["headcount"]})
                continue
            supabase.table("budget_lines").delete().eq("id", taken[key]).execute()
            taken.pop(key)
        amount = round(float(r["headcount"]) * flight, 2)
        made.append({
            "year": year,
            "department_code": department,
            "account_code": r["account_code"],
            "competition_id": r.get("competition_id"),
            "kind": "expense",
            "description": f"{r['role_label']} — {r['headcount']} × ${flight:,.0f}",
            "qty": r["headcount"],
            "monthly_amount": None,
            "amount": amount,
            "escalation_pct": 0,
            "source": "calculated",
            "notes": None,
            "series_id": None,
            "created_by": _user_id(request),
        })
    # series_id lo pone el default de la columna; mandarlo en NULL lo violaría.
    for m in made:
        m.pop("series_id")
    for i in range(0, len(made), 100):
        supabase.table("budget_lines").insert(made[i:i + 100]).execute()

    return {
        "created": len(made), "replaced": len(old),
        "avg_flight_cost": flight,
        "total": round(sum(m["amount"] for m in made), 2),
        # Salteadas para no duplicar: ya hay una línea manual en ese par.
        "skipped_conflicts": conflicts,
    }


# ─── Tarifario de fees ───────────────────────────────────────────────────────
@router.get("/fee-schedule")
def list_fee_schedule(event_type: Optional[str] = Query(None)):
    q = supabase.table("fee_schedule").select("*").eq("active", True).order("event_type")
    if event_type:
        q = q.eq("event_type", event_type)
    return q.execute().data or []


@router.put("/fee-schedule", dependencies=[Depends(require_superadmin)])
def put_fee_schedule(data: FeeScheduleUpdate):
    """Reemplaza el tarifario. Solo superadmin: es la tarifa oficial de FIBA."""
    for item in data.items:
        existing = (
            supabase.table("fee_schedule").select("id")
            .eq("role_prefix", item.role_prefix).eq("event_type", item.event_type)
            .execute().data
        )
        payload = {"fee": item.fee, "incidentals": item.incidentals,
                   "active": item.active, "notes": item.notes}
        if existing:
            supabase.table("fee_schedule").update(payload).eq("id", existing[0]["id"]).execute()
        else:
            supabase.table("fee_schedule").insert({
                "role_prefix": item.role_prefix, "event_type": item.event_type, **payload,
            }).execute()
    return {"ok": True, "count": len(data.items)}


@router.post(
    "/fee-schedule/apply",
    dependencies=[Depends(require_edit("budget")), Depends(require_edit("competitions"))],
)
def apply_fee_schedule(data: FeeScheduleApply):
    """Copia el tarifario a los defaults de fee de una competencia.

    Exige edición de `budget` **y** de `competitions`: escribe en la tabla
    `competitions`, que es de otro módulo, y esos valores terminan en las
    cartas y los pagos. Tener presupuesto no debería alcanzar para cambiarle
    el fee a un evento.

    NO crea una segunda fuente de verdad: escribe en las mismas columnas
    `{prefix}_window_fee` / `_incidentals` que `sync-nominations` ya lee (ver
    `_build_default_overrides` en games.py). Después quedan editables por
    competencia — el tarifario es el default, no una imposición.

    Por defecto no pisa lo ya cargado: si alguien ajustó el fee de un evento a
    mano, aplicar el tarifario no se lo borra. `overwrite` fuerza.

    OJO: `incidentals` del tarifario es POR CIUDAD. Para un evento multi-sede
    hay que multiplicarlo a mano — el sistema no sabe cuántas sedes tiene.
    """
    comp = supabase.table("competitions").select("*").eq("id", data.competition_id).execute().data
    if not comp:
        raise HTTPException(status_code=404, detail="Competition not found")
    comp = comp[0]

    event_type = data.fee_event_type or comp.get("fee_event_type")
    if not event_type:
        raise HTTPException(
            status_code=400,
            detail="The competition has no fee_event_type — pass one to set it",
        )

    tariffs = (
        supabase.table("fee_schedule").select("*")
        .eq("event_type", event_type).eq("active", True).execute().data
    ) or []
    if not tariffs:
        raise HTTPException(status_code=404, detail=f"No fee schedule for '{event_type}'")

    updates: dict = {}
    if event_type != comp.get("fee_event_type"):
        updates["fee_event_type"] = event_type
    applied, skipped = [], []
    for tf in tariffs:
        p = tf["role_prefix"]
        fee_col, inc_col = f"{p}_window_fee", f"{p}_incidentals"
        if comp.get(fee_col) is not None and not data.overwrite:
            skipped.append(p)
            continue
        updates[fee_col] = tf["fee"]
        updates[inc_col] = tf["incidentals"]
        applied.append({"role": p, "fee": tf["fee"], "incidentals": tf["incidentals"]})

    if updates:
        supabase.table("competitions").update(updates).eq("id", data.competition_id).execute()

    return {
        "event_type": event_type, "applied": applied, "skipped": skipped,
        # Aviso explícito: no se multiplica por sedes.
        "incidentals_are_per_city": True,
    }


@router.get("/assumptions/{year}")
def get_assumptions(year: int):
    rows = supabase.table("budget_assumptions").select("*").eq("year", year).execute().data
    return rows[0] if rows else {"year": year, "avg_flight_cost": 1050, "notes": None}


@router.put("/assumptions/{year}", dependencies=[Depends(require_edit("budget"))])
def put_assumptions(year: int, data: AssumptionsPut, request: Request):
    # El costo promedio de vuelo multiplica el headcount de Competitions: mismo
    # permiso que put_headcount.
    _assert_can_edit(request, "competitions")
    payload = {"year": year, "avg_flight_cost": data.avg_flight_cost,
               "notes": data.notes, "updated_at": _now_iso()}
    existing = supabase.table("budget_assumptions").select("year").eq("year", year).execute().data
    if existing:
        return supabase.table("budget_assumptions").update(payload).eq("year", year).execute().data[0]
    return supabase.table("budget_assumptions").insert(payload).execute().data[0]


# ─── Gastos: lectura ─────────────────────────────────────────────────────────
def _shape_expense(row: dict) -> dict:
    """Aplana los embeds de PostgREST a campos planos para la tabla del front."""
    vendor = row.pop("vendors", None) or {}
    employee = row.pop("employees", None) or {}
    comp = row.pop("competitions", None) or {}
    dept = row.pop("departments", None) or {}
    acc = row.pop("accounts", None) or {}
    row["vendor_name"] = vendor.get("name")
    row["employee_name"] = employee.get("name")
    row["competition_name"] = comp.get("name")
    row["department_label"] = dept.get("label")
    row["account_label"] = acc.get("label")
    # Nombre a mostrar, cualquiera sea el tipo de destinatario.
    row["payee_display"] = vendor.get("name") or employee.get("name") or row.get("payee_name")
    return row


_EXPENSE_SELECT = (
    "*, vendors(name), employees(name), competitions(name), "
    "departments(label), accounts(label)"
)


# `?status=pending_approval` no es un estado de la tabla: es la bandeja del
# aprobador (drafts de los departamentos donde aprueba, sin los que cargó él).
_PENDING_APPROVAL = "pending_approval"


def _fetch_expenses(
    request: Request,
    year: Optional[int] = None,
    department: Optional[str] = None,
    account: Optional[str] = None,
    competition_id: Optional[str] = None,
    general_only: bool = False,
    status: Optional[str] = None,
    vendor_id: Optional[str] = None,
    date_from: Optional[str] = None,
    date_to: Optional[str] = None,
    search: Optional[str] = None,
    scope_action: str = "view",
    created_by: Optional[str] = None,
    exclude_created_by: Optional[str] = None,
) -> list:
    """Compartido por /expenses y /expenses/summary.

    Los endpoints no se llaman entre sí como funciones: los defaults `Query()`
    se filtran como objetos truthy y se vuelven filtros fantasma (mismo bug que
    dejó el summary de payments en $0.00).

    `scope_action='approve'` recorta a los departamentos donde el caller ve Y
    aprueba; lo resuelve _scoped() en un único filtro sobre department_code.
    """
    q = _scoped(
        supabase.table("expenses").select(_EXPENSE_SELECT).order("expense_date", desc=True),
        request, department, action=scope_action,
    )

    if account:
        q = q.eq("account_code", account)
    if status:
        q = q.eq("status", status)
    if vendor_id:
        q = q.eq("vendor_id", vendor_id)
    if created_by:
        q = q.eq("created_by", created_by)
    if general_only:
        q = q.is_("competition_id")
    elif competition_id:
        q = q.eq("competition_id", competition_id)

    # `year` es azúcar sobre el rango; si vienen los dos, gana el rango explícito.
    lower = date_from or (f"{year}-01-01" if year else None)
    upper = date_to or (f"{year}-12-31" if year else None)
    # ⚠️ El builder guarda UN filtro por columna: un .gte() y un .lte() sobre
    # expense_date se pisan, y "año 2026" devolvía todo lo anterior a 2027.
    # La cota inferior va a la query; si además hay superior, se aplica en
    # Python sobre lo que ya vino recortado.
    if lower:
        q = q.gte("expense_date", lower)
    elif upper:
        q = q.lte("expense_date", upper)

    rows = [_shape_expense(r) for r in (q.execute().data or [])]

    if lower and upper:
        rows = [r for r in rows if str(r.get("expense_date") or "")[:10] <= upper[:10]]
    # `neq` en PostgREST también descarta los NULL (los 129 importados no
    # tienen created_by), por eso la exclusión del propio va en Python.
    if exclude_created_by:
        rows = [r for r in rows if r.get("created_by") != exclude_created_by]

    if search:
        s = search.lower()
        rows = [
            r for r in rows
            if s in (r.get("description") or "").lower()
            or s in (r.get("record_no") or "").lower()
            or s in (r.get("payee_display") or "").lower()
            or s in (r.get("invoice_no") or "").lower()
        ]
    return rows


@router.get("/expenses")
def list_expenses(
    request: Request,
    year: Optional[int] = Query(None),
    department: Optional[str] = Query(None),
    account: Optional[str] = Query(None),
    competition_id: Optional[str] = Query(None),
    general_only: bool = Query(False, description="Solo gastos sin competencia"),
    status: Optional[str] = Query(
        None, description="Estado, o 'pending_approval' = drafts que el caller puede aprobar",
    ),
    vendor_id: Optional[str] = Query(None),
    date_from: Optional[str] = Query(None),
    date_to: Optional[str] = Query(None),
    search: Optional[str] = Query(None),
    mine: bool = Query(False, description="Solo los que cargó el caller"),
):
    return _expense_rows(
        request, year, department, account, competition_id, general_only,
        status, vendor_id, date_from, date_to, search, mine,
    )


def _expense_rows(
    request: Request,
    year: Optional[int],
    department: Optional[str],
    account: Optional[str],
    competition_id: Optional[str],
    general_only: bool,
    status: Optional[str],
    vendor_id: Optional[str],
    date_from: Optional[str],
    date_to: Optional[str],
    search: Optional[str],
    mine: bool,
) -> list:
    """Cuerpo de GET /expenses, compartido con /expenses/export.xlsx para que
    el Excel traiga exactamente las filas de la tabla con los mismos filtros
    (incluidos `pending_approval` y `mine`). Todo pasa por _fetch_expenses →
    _scoped(): el export no ve nada que la pantalla no muestre."""
    scope_action, exclude = "view", None
    if status == _PENDING_APPROVAL:
        # Drafts de los departamentos donde aprueba, sin los propios: nadie
        # aprueba lo que cargó (el superadmin puede con override, pero desde el
        # gasto, no desde la bandeja). Es el mismo criterio que el contador de
        # la campana (notifications.py, /pending).
        status, scope_action, exclude = "draft", "approve", _user_id(request)
    elif status and status not in _EXPENSE_STATUSES:
        raise HTTPException(status_code=400, detail="Invalid status")

    created_by = None
    if mine:
        created_by = _user_id(request)
        if not created_by:
            return []

    return _fetch_expenses(
        request, year, department, account, competition_id, general_only,
        status, vendor_id, date_from, date_to, search,
        scope_action=scope_action, created_by=created_by, exclude_created_by=exclude,
    )


# Declarado ANTES de /expenses/{expense_id} — si no, "summary" entra como id.
@router.get("/expenses/summary")
def expenses_summary(
    request: Request,
    year: Optional[int] = Query(None),
    department: Optional[str] = Query(None),
    account: Optional[str] = Query(None),
    competition_id: Optional[str] = Query(None),
    general_only: bool = Query(False),
    date_from: Optional[str] = Query(None),
    date_to: Optional[str] = Query(None),
):
    rows = _fetch_expenses(
        request, year=year, department=department, account=account,
        competition_id=competition_id, general_only=general_only,
        date_from=date_from, date_to=date_to,
    )

    def _total(predicate) -> float:
        return round(sum(float(r.get("amount") or 0) for r in rows if predicate(r)), 2)

    live = [r for r in rows if r.get("status") not in ("rejected", "cancelled")]
    return {
        "count": len(rows),
        # Ejecutado: lo único que consume presupuesto.
        "paid": _total(lambda r: r.get("status") == "paid"),
        # Comprometido: aprobado y todavía sin pagar.
        "committed": _total(lambda r: r.get("status") == "approved"),
        "draft": _total(lambda r: r.get("status") == "draft"),
        "total": round(sum(float(r.get("amount") or 0) for r in live), 2),
        "by_department": _group_totals(live, "department_code"),
        "by_account": _group_totals(live, "account_code"),
    }


_EXPENSE_EXPORT_HEADERS = {
    "es": ["Registro", "Fecha", "Departamento", "Cuenta", "Competencia", "Descripción",
           "Destinatario", "Monto", "Estado", "Fecha de pago", "Método",
           "Confirmación bancaria", "Factura", "Línea de presupuesto", "Creado"],
    "en": ["Record", "Date", "Department", "Account", "Competition", "Description",
           "Payee", "Amount", "Status", "Payment date", "Method",
           "Bank confirmation", "Invoice", "Budget line", "Created"],
}
_EXPENSE_EXPORT_AMOUNT_COL = 8   # 1-based, "Monto"
_EXPENSE_STATUS_LABELS = {
    "es": {"draft": "Borrador", "approved": "Aprobado", "paid": "Pagado",
           "rejected": "Rechazado", "cancelled": "Cancelado"},
    "en": {"draft": "Draft", "approved": "Approved", "paid": "Paid",
           "rejected": "Rejected", "cancelled": "Cancelled"},
}
# Ids por request en el `in.(...)` de PostgREST: 100 UUIDs entran holgados en
# la URL; más que eso se parte en varias queries (sigue sin ser N+1).
_IN_CHUNK = 100


def _line_descriptions(request: Request, line_ids: list) -> dict:
    """{budget_line_id: description} para las líneas que referencian los gastos.

    Pasa por _scoped() igual que cualquier query de budget: una línea de un
    departamento que el caller no ve sale en blanco en vez de filtrar su
    descripción (no debería pasar, _check_budget_line exige el mismo depto).
    """
    out: dict = {}
    for i in range(0, len(line_ids), _IN_CHUNK):
        chunk = line_ids[i:i + _IN_CHUNK]
        q = _scoped(
            supabase.table("budget_lines").select("id, description").in_("id", chunk),
            request,
        )
        out.update({r["id"]: r.get("description") for r in (q.execute().data or [])})
    return out


def _bold_last_row(ws) -> None:
    from openpyxl.styles import Font

    for cell in ws[ws.max_row]:
        cell.font = Font(bold=True)


# Declarado ANTES de cualquier /expenses/{expense_id}: si no, "export.xlsx"
# entraría como id.
@router.get("/expenses/export.xlsx")
def export_expenses(
    request: Request,
    year: Optional[int] = Query(None),
    department: Optional[str] = Query(None),
    account: Optional[str] = Query(None),
    competition_id: Optional[str] = Query(None),
    general_only: bool = Query(False),
    status: Optional[str] = Query(None),
    vendor_id: Optional[str] = Query(None),
    date_from: Optional[str] = Query(None),
    date_to: Optional[str] = Query(None),
    search: Optional[str] = Query(None),
    mine: bool = Query(False),
    lang: str = Query("es"),
):
    """Los gastos de la tabla, con sus mismos filtros, en Excel.

    Los nombres (departamento, cuenta, competencia, destinatario) ya vienen en
    los embeds de _EXPENSE_SELECT; las líneas de presupuesto se resuelven con
    un solo diccionario. Nada de una query por fila.
    """
    from openpyxl import Workbook
    from api._lib.services.xlsx_utils import (
        autosize, money_format, safe_text, style_header, workbook_bytes, xlsx_response,
    )

    lang_key = "en" if lang == "en" else "es"
    rows = _expense_rows(
        request, year, department, account, competition_id, general_only,
        status, vendor_id, date_from, date_to, search, mine,
    )
    lines = _line_descriptions(
        request, sorted({r["budget_line_id"] for r in rows if r.get("budget_line_id")}),
    )
    status_labels = _EXPENSE_STATUS_LABELS[lang_key]
    general = "General"

    wb = Workbook()
    ws = wb.active
    ws.title = "Expenses" if lang_key == "en" else "Gastos"
    ws.append(_EXPENSE_EXPORT_HEADERS[lang_key])
    style_header(ws)

    total = 0.0
    for r in rows:
        amount = float(r.get("amount") or 0)
        total += amount
        acc_code, acc_label = r.get("account_code"), r.get("account_label")
        ws.append([
            safe_text(r.get("record_no")),
            str(r.get("expense_date") or "")[:10] or None,
            safe_text(r.get("department_label") or r.get("department_code")),
            safe_text(f"{acc_code} — {acc_label}" if acc_label else acc_code),
            safe_text(r.get("competition_name") or (general if not r.get("competition_id") else None)),
            safe_text(r.get("description")),
            safe_text(r.get("payee_display")),
            amount,
            status_labels.get(r.get("status"), r.get("status")),
            str(r.get("payment_date") or "")[:10] or None,
            safe_text(r.get("payment_method")),
            safe_text(r.get("bank_confirmation")),
            safe_text(r.get("invoice_no")),
            safe_text(lines.get(r.get("budget_line_id"))) if r.get("budget_line_id") else None,
            str(r.get("created_at") or "")[:16].replace("T", " ") or None,
        ])

    # Total de la columna Monto tal como se exportó: si el filtro trae
    # rechazados o cancelados, suman acá (es la suma de lo que se ve).
    footer = [None] * len(_EXPENSE_EXPORT_HEADERS[lang_key])
    footer[0] = "Total"
    footer[_EXPENSE_EXPORT_AMOUNT_COL - 1] = round(total, 2)
    ws.append(footer)
    _bold_last_row(ws)

    money_format(ws, [_EXPENSE_EXPORT_AMOUNT_COL])
    autosize(ws)
    return xlsx_response(workbook_bytes(wb), f"budget-expenses-{year or 'all'}.xlsx")


def _group_totals(rows: list, key: str) -> list:
    acc: dict = {}
    for r in rows:
        k = r.get(key) or "—"
        entry = acc.setdefault(k, {"key": k, "paid": 0.0, "committed": 0.0, "total": 0.0})
        amount = float(r.get("amount") or 0)
        entry["total"] += amount
        if r.get("status") == "paid":
            entry["paid"] += amount
        elif r.get("status") == "approved":
            entry["committed"] += amount
    for e in acc.values():
        for f in ("paid", "committed", "total"):
            e[f] = round(e[f], 2)
    return sorted(acc.values(), key=lambda e: -e["total"])


# ─── Gastos: máquina de estados ──────────────────────────────────────────────
#
# El `status` de un gasto NO se toca por PATCH. Solo por
# POST /expenses/{id}/transition, que pasa por check_expense_transition().
#
# Por qué: con el status editable en el PATCH, el mismo designado que cargaba un
# gasto lo aprobaba y lo marcaba pagado en un solo formulario, sin motivo ni
# traza — un único sello. La máquina de estados es lo que hace cumplir la
# separación de funciones (quien carga no aprueba), exige el motivo donde hace
# falta (rechazo, cancelación, reversión de un pago) y deja cada cambio en
# finance_status_history. Si el status volviera al PATCH, todo eso se saltea
# con un `{"status": "paid"}`.
#
#   draft     → approved   aprobador del depto, created_by ≠ caller
#                          (superadmin sobre lo propio: reason + meta.override)
#   draft     → rejected   aprobador del depto, reason
#   draft     → cancelled  editor del depto, reason
#   approved  → paid       editor del depto, payment_date (sella paid_by/at)
#   approved  → draft      aprobador del depto o superadmin (limpia approved_*)
#   rejected  → draft      editor del depto (limpia los sellos)
#   cancelled → draft      editor del depto (limpia los sellos)
#   paid      → approved   solo superadmin, reason (limpia paid_* y payment_date)
#
# Cualquier otra combinación es 409. No hay approved → cancelled: un gasto
# aprobado se reabre (→ draft) y recién ahí se cancela, así la cancelación la
# decide el editor sobre un borrador y no se pierde el rastro de la aprobación.
_EXPENSE_TRANSITIONS = {
    ("draft", "approved"),
    ("draft", "rejected"),
    ("draft", "cancelled"),
    ("approved", "paid"),
    ("approved", "draft"),
    ("rejected", "draft"),
    ("cancelled", "draft"),
    ("paid", "approved"),
}

# Volver a borrador borra los sellos de la salida anterior: un draft no tiene
# aprobador, ni rechazo, ni cancelación vigentes (el historial sí los guarda).
_EXPENSE_STAMPS_CLEARED = {
    "approved_by": None, "approved_at": None,
    "rejected_by": None, "rejected_at": None, "rejection_reason": None,
    "cancelled_by": None, "cancelled_at": None, "cancel_reason": None,
}

# Lo que define QUÉ plata se comprometió: a quién, cuánto, de qué departamento,
# a qué cuenta, por qué evento y con qué fecha. Aprobado eso, cambiarlo sin
# volver a draft dejaría una aprobación sobre algo que el aprobador no vio.
_EXPENSE_FINANCIAL_FIELDS = (
    "department_code", "account_code", "competition_id", "amount", "expense_date",
    "payee_type", "vendor_id", "employee_id", "payee_name",
)
# Un gasto pagado es un hecho contable: solo se le completan datos administrativos.
_EXPENSE_PAID_EDITABLE = {"comments", "invoice_no"}
_EXPENSE_DELETABLE = ("draft", "rejected", "cancelled")


def _clean_text(value) -> Optional[str]:
    return (str(value).strip() or None) if value is not None else None


def _iso_date_or_400(value, field: str) -> str:
    try:
        return date.fromisoformat(str(value)[:10]).isoformat()
    except (TypeError, ValueError):
        raise HTTPException(status_code=400, detail=f"{field} must be a date (YYYY-MM-DD)")


def check_expense_transition(
    expense: dict,
    target: str,
    payload=None,
    *,
    user_id: Optional[str],
    can_edit: bool,
    can_approve: bool,
    is_superadmin: bool,
    now: Optional[str] = None,
) -> dict:
    """Valida una transición de estado de un gasto. Pura: no toca la base.

    `expense` es la fila actual; `target`, el estado destino; `payload`, un dict
    (o el ExpenseTransition) con reason / payment_date / payment_method /
    bank_confirmation. `can_edit` / `can_approve` son los derechos del caller
    SOBRE EL DEPARTAMENTO DEL GASTO, ya resueltos por quien llama.

    Devuelve {from_status, to_status, updates, reason, meta}: `updates` son las
    columnas a escribir (status incluido, sin updated_at). Levanta
    HTTPException 400 (falta un dato), 403 (no le corresponde al caller) o 409
    (la transición no existe desde el estado actual).
    """
    if payload is None:
        payload = {}
    elif hasattr(payload, "model_dump"):
        payload = payload.model_dump()
    now = now or _now_iso()
    current = expense.get("status")
    dept = expense.get("department_code")

    if target not in _EXPENSE_STATUSES:
        raise HTTPException(status_code=400, detail="Invalid status")
    if target == current:
        raise HTTPException(status_code=409, detail=f"The expense is already {target}")
    if (current, target) not in _EXPENSE_TRANSITIONS:
        raise HTTPException(status_code=409, detail=f"An expense cannot go from {current} to {target}")

    reason = _clean_text(payload.get("reason"))
    updates: dict = {"status": target}
    meta: dict = {}

    def need_reason(what: str) -> None:
        if not reason:
            raise HTTPException(status_code=400, detail=f"A reason is required to {what}")

    def need(allowed: bool, detail: str) -> None:
        if not allowed:
            raise HTTPException(status_code=403, detail=detail)

    if (current, target) == ("draft", "approved"):
        need(can_approve, f"No approval rights on department '{dept}'")
        creator = expense.get("created_by")
        if creator and creator == user_id:
            # Separación de funciones. El superadmin puede saltearla (Finance es
            # una persona en la práctica), pero con motivo y marcado como
            # override en el historial: el control es la traza, no el bloqueo.
            if not is_superadmin:
                raise HTTPException(
                    status_code=403,
                    detail="You cannot approve an expense you created: another approver has to",
                )
            need_reason("approve your own expense (it is recorded as an override)")
            meta["override"] = True
        updates.update(approved_by=user_id, approved_at=now)

    elif (current, target) == ("draft", "rejected"):
        need(can_approve, f"No approval rights on department '{dept}'")
        need_reason("reject an expense")
        updates.update(rejected_by=user_id, rejected_at=now, rejection_reason=reason)

    elif (current, target) == ("draft", "cancelled"):
        need(can_edit, f"No edit access to department '{dept}'")
        need_reason("cancel an expense")
        updates.update(cancelled_by=user_id, cancelled_at=now, cancel_reason=reason)

    elif (current, target) == ("approved", "paid"):
        need(can_edit, f"No edit access to department '{dept}'")
        if not payload.get("payment_date"):
            raise HTTPException(status_code=400, detail="A payment date is required to mark an expense as paid")
        payment_date = _iso_date_or_400(payload.get("payment_date"), "payment_date")
        updates.update(payment_date=payment_date, paid_by=user_id, paid_at=now)
        for field in ("payment_method", "bank_confirmation"):
            value = _clean_text(payload.get(field))
            if value:
                updates[field] = value
        meta["payment_date"] = payment_date

    elif (current, target) == ("approved", "draft"):
        need(can_approve or is_superadmin, f"No approval rights on department '{dept}'")
        updates.update(approved_by=None, approved_at=None)

    elif target == "draft":   # rejected / cancelled → draft
        need(can_edit, f"No edit access to department '{dept}'")
        updates.update(_EXPENSE_STAMPS_CLEARED)

    elif (current, target) == ("paid", "approved"):
        # Revertir un pago saca plata del ejecutado: solo superadmin y con motivo.
        need(is_superadmin, "Only a superadmin can revert a paid expense")
        need_reason("revert a paid expense")
        updates.update(payment_date=None, paid_by=None, paid_at=None)
        meta["reverted_payment_date"] = expense.get("payment_date")

    return {
        "from_status": current,
        "to_status": target,
        "updates": updates,
        "reason": reason,
        "meta": meta or None,
    }


def _expense_title(expense: dict) -> str:
    """'GA-00042 · IT · $1,250.00' — título de la novedad en la campana."""
    code = expense.get("department_code") or "—"
    label = code
    try:
        rows = supabase.table("departments").select("label").eq("code", code).execute().data
        if rows and rows[0].get("label"):
            label = rows[0]["label"]
    except Exception:
        pass
    try:
        amount = f"${float(expense.get('amount') or 0):,.2f}"
    except (TypeError, ValueError):
        amount = f"${expense.get('amount')}"
    return f"{expense.get('record_no') or 'GA'} · {label} · {amount}"


def _expense_link(expense_id) -> str:
    return f"/budget?tab=expenses&open={expense_id}"


def _notify_expense_submitted(expense: dict, request: Request) -> None:
    """Un draft nuevo (alta o vuelta a borrador) → aprobadores del depto.

    Sin el creador (no aprueba lo propio) ni quien hizo la acción.
    """
    notify(
        budget_approvers_for(expense.get("department_code")),
        "expense_submitted",
        _expense_title(expense),
        body=expense.get("description"),
        entity_type="expense", entity_id=expense.get("id"),
        link=_expense_link(expense.get("id")),
        exclude=[expense.get("created_by"), _user_id(request)],
    )


_EXPENSE_NOTIFY_CREATOR = {
    "approved": "expense_approved",
    "rejected": "expense_rejected",
    "paid": "expense_paid",
}


def _notify_expense_transition(expense: dict, plan: dict, request: Request) -> None:
    """Novedad de la transición. Best-effort: notify() nunca levanta."""
    try:
        to_status, from_status = plan["to_status"], plan["from_status"]
        if to_status == "draft":
            _notify_expense_submitted(expense, request)
            return
        kind = _EXPENSE_NOTIFY_CREATOR.get(to_status)
        # paid → approved es una reversión, no "tu gasto fue aprobado".
        if not kind or (to_status == "approved" and from_status != "draft"):
            return
        body = expense.get("description")
        if to_status == "rejected" and plan.get("reason"):
            body = plan["reason"]
        notify(
            expense.get("created_by"),
            kind,
            _expense_title(expense),
            body=body,
            entity_type="expense", entity_id=expense.get("id"),
            link=_expense_link(expense.get("id")),
            exclude=_user_id(request),
        )
    except Exception as exc:
        logger.warning("[budget] expense notification failed: %s", exc)


def _check_budget_line(request: Request, line_id: str, department_code: str,
                       account_code: str, expense_date) -> None:
    """La línea de presupuesto a la que se imputa el gasto tiene que ser la suya.

    Mismo departamento, misma cuenta y el año fiscal de expense_date. Sin esto
    el gasto de IT podría "consumir" una línea de Comms o de otro año, y el
    ejecutado por línea quedaría torcido sin que nada falle.
    """
    try:
        uuid.UUID(str(line_id))
    except ValueError:
        raise HTTPException(status_code=400, detail="Unknown budget line")
    rows = (
        supabase.table("budget_lines")
        .select("id, year, department_code, account_code, kind")
        .eq("id", line_id)
        .execute()
        .data
    )
    # Una línea de un departamento que el caller no ve responde igual que una
    # inexistente: un mensaje distinto confirmaría que existe.
    if not rows or not _can_view_department(request, rows[0].get("department_code")):
        raise HTTPException(status_code=400, detail="Unknown budget line")
    line = rows[0]
    if line.get("kind") != "expense":
        raise HTTPException(status_code=400, detail="The budget line is not an expense line")
    if line.get("department_code") != department_code or line.get("account_code") != account_code:
        raise HTTPException(
            status_code=400,
            detail="The budget line must have the same department and account as the expense",
        )
    if str(line.get("year")) != str(expense_date or "")[:4]:
        raise HTTPException(
            status_code=400,
            detail=f"The budget line is for {line.get('year')}; the expense date is in another year",
        )


def _same_value(field: str, new, old) -> bool:
    """¿El PATCH deja el campo como estaba? (el editor manda el form entero)."""
    if new in (None, "") and old in (None, ""):
        return True
    if new is None or old is None:
        return False
    if field == "amount":
        try:
            return round(float(new), 2) == round(float(old), 2)
        except (TypeError, ValueError):
            return False
    if field in ("expense_date", "payment_date"):
        return str(new)[:10] == str(old)[:10]
    return str(new) == str(old)


# ─── Gastos: escritura ───────────────────────────────────────────────────────
@router.post("/expenses", status_code=201, dependencies=[Depends(require_edit("budget"))])
def create_expense(data: ExpenseCreate, request: Request):
    if not _valid_department(data.department_code):
        raise HTTPException(status_code=400, detail="Unknown or inactive department")
    _assert_can_edit(request, data.department_code)
    _check_account(data.account_code, "expense")

    payee_type = data.payee_type or "vendor"
    _check_payee(payee_type, data.vendor_id, data.employee_id, data.payee_name)

    # Un gasto nace draft y lo aprueba OTRO por /transition. `approved` directo
    # sería el auto-sello que la máquina de estados viene a cerrar. `paid`
    # directo queda solo para la carga histórica (gastos que ya salieron, como
    # los 129 de IT 2026) y solo para el admin general.
    status = data.status or "draft"
    if status == "approved":
        raise HTTPException(
            status_code=400,
            detail="An expense is created as draft; it is approved by an approver via /transition",
        )
    if status not in ("draft", "paid"):
        raise HTTPException(status_code=400, detail="A new expense must be 'draft' (or 'paid' for a historical load)")
    if status == "paid" and not _is_general_admin(request):
        raise HTTPException(status_code=403, detail="Only a general budget admin can load an already-paid expense")
    _check_paid_coherence(status, data.payment_date)

    if data.amount is None or float(data.amount) < 0:
        raise HTTPException(status_code=400, detail="amount must be zero or positive")
    if data.budget_line_id:
        _check_budget_line(request, data.budget_line_id, data.department_code,
                           data.account_code, data.expense_date)

    record = {k: v for k, v in data.model_dump().items() if v is not None}
    record["payee_type"] = payee_type
    record["status"] = status
    record["created_by"] = _user_id(request)
    # Los campos del destinatario que no corresponden al tipo se limpian: si no,
    # un gasto que pasó de vendor a employee queda con las dos FK cargadas y
    # payee_display elige la equivocada.
    if payee_type != "vendor":
        record.pop("vendor_id", None)
    if payee_type != "employee":
        record.pop("employee_id", None)

    if status == "paid":
        record["payment_date"] = _iso_date_or_400(data.payment_date, "payment_date")
        record["paid_by"] = _user_id(request)
        record["paid_at"] = _now_iso()

    row = supabase.table("expenses").insert(record).execute().data[0]
    record_transition(
        "expense", row["id"], None, status, request,
        meta={"historical_load": True, "payment_date": record["payment_date"]} if status == "paid" else None,
    )
    if status == "draft":
        _notify_expense_submitted(row, request)
    return row


@router.patch("/expenses/{expense_id}", dependencies=[Depends(require_edit("budget"))])
def update_expense(expense_id: str, data: ExpenseUpdate, request: Request):
    current = _row_or_403("expenses", expense_id, request, "edit")
    current_status = current.get("status")

    updates = _patch_updates(data, "expenses")

    # status y payment_date SOLO por /transition (ver la cabecera de la máquina
    # de estados). Si vienen con el valor que ya tenían se ignoran — el editor
    # del front manda el formulario entero —; si cambian, 400.
    for field in ("status", "payment_date"):
        if field in updates:
            if not _same_value(field, updates[field], current.get(field)):
                raise HTTPException(
                    status_code=400,
                    detail=f"{field} cannot be changed here: use POST /budget/expenses/{{id}}/transition",
                )
            updates.pop(field)

    changed = {f for f, v in updates.items() if not _same_value(f, v, current.get(f))}
    if current_status == "approved":
        locked = sorted(changed & set(_EXPENSE_FINANCIAL_FIELDS))
        if locked:
            raise HTTPException(
                status_code=409,
                detail=f"The expense is approved: reopen it (back to draft) before changing {', '.join(locked)}",
            )
    elif current_status == "paid":
        locked = sorted(changed - _EXPENSE_PAID_EDITABLE)
        if locked:
            raise HTTPException(
                status_code=409,
                detail=f"The expense is paid: only comments and invoice_no can change (not {', '.join(locked)})",
            )
    if not updates:
        raise HTTPException(status_code=400, detail="Nothing to update")

    if "department_code" in updates:
        if not _valid_department(updates["department_code"]):
            raise HTTPException(status_code=400, detail="Unknown or inactive department")
        _assert_can_edit(request, updates["department_code"])
    if "account_code" in updates:
        _check_account(updates["account_code"], "expense")
    if "amount" in updates and float(updates["amount"]) < 0:
        raise HTTPException(status_code=400, detail="amount must be zero or positive")

    # El destinatario se revalida solo si el PATCH lo toca: un gasto pagado
    # importado con datos viejos no puede quedar trabado para editar un comentario.
    if {"payee_type", "vendor_id", "employee_id", "payee_name"} & updates.keys():
        payee_type = updates.get("payee_type", current.get("payee_type"))
        _check_payee(
            payee_type,
            updates.get("vendor_id", current.get("vendor_id")),
            updates.get("employee_id", current.get("employee_id")),
            updates.get("payee_name", current.get("payee_name")),
        )
        if payee_type != "vendor":
            updates["vendor_id"] = None
        if payee_type != "employee":
            updates["employee_id"] = None

    # La línea de presupuesto tiene que seguir siendo del mismo depto/cuenta/año,
    # tanto si la cambian como si cambian lo que la ata (depto, cuenta, fecha).
    line_id = updates.get("budget_line_id", current.get("budget_line_id"))
    if line_id and ({"budget_line_id", "department_code", "account_code", "expense_date"} & changed):
        _check_budget_line(
            request, line_id,
            updates.get("department_code", current.get("department_code")),
            updates.get("account_code", current.get("account_code")),
            updates.get("expense_date", current.get("expense_date")),
        )

    updates["updated_at"] = _now_iso()
    # El UPDATE exige el estado que se validó: si en el medio alguien lo aprobó,
    # no se le cambian los campos financieros a un gasto ya aprobado.
    result = (
        supabase.table("expenses").update(updates)
        .eq("id", expense_id).eq("status", current_status)
        .execute()
    )
    if not result.data:
        raise HTTPException(status_code=409, detail="The expense changed or was removed meanwhile; reload it")
    return result.data[0]


def _run_expense_transition(expense_id: str, data: ExpenseTransition, request: Request) -> dict:
    """Aplica una transición: permisos → check puro → UPDATE → historial → aviso."""
    # Se trae con "view" y no con "edit": quién puede depende de la transición
    # (aprobar es de can_approve, que no implica can_edit). Si no ve el
    # departamento, 404 igual que inexistente.
    current = _row_or_403("expenses", expense_id, request, "view")
    dept = current.get("department_code")
    edit_depts = _departments_for(request, "edit")
    approve_depts = _departments_for(request, "approve")
    # "Editor" = el permiso del módulo (lo que exige require_edit en el resto de
    # las escrituras) MÁS can_edit sobre el departamento en budget_access.
    can_edit = has_edit(request, "budget") and (edit_depts is None or dept in edit_depts)
    can_approve = approve_depts is None or dept in approve_depts

    plan = check_expense_transition(
        current, (data.status or "").strip(), data,
        user_id=_user_id(request),
        can_edit=can_edit,
        can_approve=can_approve,
        is_superadmin=is_superadmin_request(request),
    )
    updates = {**plan["updates"], "updated_at": _now_iso()}
    # Concurrencia: el UPDATE exige que el estado siga siendo el validado. Dos
    # aprobadores que aprietan a la vez → uno gana, el otro recibe 409.
    result = (
        supabase.table("expenses").update(updates)
        .eq("id", expense_id).eq("status", plan["from_status"])
        .execute()
    )
    if not result.data:
        raise HTTPException(status_code=409, detail="The expense changed meanwhile; reload it and try again")
    row = result.data[0]
    record_transition(
        "expense", expense_id, plan["from_status"], plan["to_status"], request,
        reason=plan["reason"], meta=plan["meta"],
    )
    _notify_expense_transition(row, plan, request)
    return row


# Sin require_edit en el decorador, A PROPÓSITO: aprobar/rechazar es de
# `budget_access.can_approve`, que no implica el can_edit del módulo (el que
# aprueba no tiene por qué poder cargar). El router ya exige require_view
# ("budget"), y check_expense_transition() exige, por transición, can_approve
# del departamento o "editor" (= has_edit("budget") + can_edit del depto), o
# superadmin. No hay camino que escriba sin uno de esos tres.
@router.post("/expenses/{expense_id}/transition")
def transition_expense(expense_id: str, data: ExpenseTransition, request: Request):
    return _run_expense_transition(expense_id, data, request)


@router.post("/expenses/{expense_id}/approve")
def approve_expense(expense_id: str, request: Request, data: Optional[ExpenseApprove] = None):
    """Alias de /transition con status=approved (mismos permisos y reglas).

    Queda por compatibilidad con el front anterior. El body es opcional; solo
    hace falta `reason` cuando un superadmin aprueba un gasto propio.
    """
    return _run_expense_transition(
        expense_id,
        ExpenseTransition(status="approved", reason=data.reason if data else None),
        request,
    )


@router.get("/expenses/{expense_id}/history")
def expense_history(expense_id: str, request: Request):
    """Timeline de estados del gasto (finance_status_history), más reciente primero.

    El scope se valida contra el gasto: sin acceso a su departamento, 404.
    """
    _row_or_403("expenses", expense_id, request)
    return history_for("expense", expense_id)


@router.delete("/expenses/{expense_id}", dependencies=[Depends(require_edit("budget"))])
def delete_expense(expense_id: str, request: Request):
    """Borra el gasto y limpia sus adjuntos de Storage.

    Solo draft / rejected / cancelled: un gasto aprobado o pagado ya es parte
    del comprometido o del ejecutado, y borrarlo lo haría desaparecer del
    reporte sin traza. Se cancela con motivo (si está aprobado, reabriéndolo
    primero); un pagado solo lo revierte un superadmin.

    El CASCADE de la FK borra las filas de expense_attachments, pero no los
    objetos del bucket: hay que sacarlos a mano (mismo patrón que payments).
    """
    current = _row_or_403("expenses", expense_id, request, "edit")
    if current.get("status") not in _EXPENSE_DELETABLE:
        raise HTTPException(
            status_code=409,
            detail="Only draft, rejected or cancelled expenses can be deleted; cancel it instead",
        )
    atts = (
        supabase.table("expense_attachments")
        .select("storage_path")
        .eq("expense_id", expense_id)
        .execute()
        .data
    ) or []
    # El filtro de estado va también en el DELETE: si entre la lectura y acá lo
    # aprobaron, no se borra.
    result = (
        supabase.table("expenses").delete()
        .eq("id", expense_id).in_("status", list(_EXPENSE_DELETABLE))
        .execute()
    )
    if not result.data:
        raise HTTPException(status_code=409, detail="The expense changed or was removed meanwhile; reload it")
    for a in atts:
        _delete_from_storage(a.get("storage_path"))
    return {"ok": True}


# ─── Adjuntos ────────────────────────────────────────────────────────────────
# /attachments/... va antes que /expenses/{id}/... por claridad de rutas; no hay
# choque porque los prefijos difieren, pero mantenerlos juntos evita sorpresas.
def _attachment_or_403(attachment_id: str, request: Request, action: str = "view") -> dict:
    """El adjunto hereda el departamento de su gasto.

    Sin esto el scoping tendría un agujero por la puerta de atrás: los adjuntos
    se piden por su propio id, así que un designado podría bajarse la factura de
    otro departamento adivinando (o filtrando) un uuid.
    """
    rows = (
        supabase.table("expense_attachments")
        .select("*, expenses(department_code)")
        .eq("id", attachment_id)
        .execute()
        .data
    )
    if not rows:
        raise HTTPException(status_code=404, detail="Attachment not found")
    row = rows[0]
    dept = (row.pop("expenses", None) or {}).get("department_code")
    if not _can_view_department(request, dept):
        raise HTTPException(status_code=404, detail="Attachment not found")
    if action == "edit":
        _assert_can_edit(request, dept)
    return row


@router.get("/attachments/{attachment_id}/download")
def download_attachment(attachment_id: str, request: Request, filename: Optional[str] = None):
    """Descarga autenticada por blob. El bucket es privado: nunca URL pública."""
    row = [_attachment_or_403(attachment_id, request)]

    key = _extract_storage_key(row[0].get("storage_path"))
    if not key:
        raise HTTPException(status_code=404, detail="File not available")

    filename = filename or row[0].get("file_name") or "attachment"
    filename = _SAFE_FILENAME_RE.sub("", filename).strip() or "attachment"

    supabase_url = os.environ.get("SUPABASE_URL", "")
    supabase_key = os.environ.get("SUPABASE_SERVICE_ROLE_KEY", "") or os.environ.get("SUPABASE_KEY", "")
    object_url = f"{supabase_url}/storage/v1/object/{_BUCKET}/{key}"
    try:
        headers = {"Authorization": f"Bearer {supabase_key}", "apikey": supabase_key}
        resp = httpx.get(object_url, headers=headers, timeout=30.0, follow_redirects=True)
        if resp.status_code != 200:
            raise HTTPException(status_code=404, detail="File not available")
        return Response(
            content=resp.content,
            media_type=resp.headers.get("content-type", "application/octet-stream"),
            headers={
                "Content-Disposition": f'attachment; filename="{filename}"',
                "Cache-Control": "private, no-store",
            },
        )
    except HTTPException:
        raise
    except Exception:
        raise HTTPException(status_code=502, detail="Storage fetch failed")


@router.delete("/attachments/{attachment_id}", dependencies=[Depends(require_edit("budget"))])
def delete_attachment(attachment_id: str, request: Request):
    row = _attachment_or_403(attachment_id, request, "edit")
    supabase.table("expense_attachments").delete().eq("id", attachment_id).execute()
    _delete_from_storage(row.get("storage_path"))
    return {"ok": True}


@router.get("/expenses/{expense_id}/attachments")
def list_attachments(expense_id: str, request: Request):
    _row_or_403("expenses", expense_id, request)
    return (
        supabase.table("expense_attachments")
        .select("id, file_name, kind, uploaded_at")
        .eq("expense_id", expense_id)
        .order("uploaded_at")
        .execute()
        .data
    ) or []


@router.post(
    "/expenses/{expense_id}/attachments",
    status_code=201,
    dependencies=[Depends(require_edit("budget"))],
)
async def upload_attachment(
    expense_id: str,
    request: Request,
    file: UploadFile = File(...),
    kind: Optional[str] = Form(None),
):
    _row_or_403("expenses", expense_id, request, "edit")

    content_type = file.content_type or "application/octet-stream"
    if not any(content_type.startswith(a) for a in _ALLOWED_CONTENT):
        raise HTTPException(status_code=400, detail="Only PDF or image files are allowed")

    content = await file.read()
    if len(content) > _MAX_UPLOAD_BYTES:
        raise HTTPException(status_code=400, detail="File too large (max 10 MB)")

    ext = (file.filename or "").rsplit(".", 1)[-1].lower() if "." in (file.filename or "") else "bin"
    ext = re.sub(r"[^a-z0-9]", "", ext)[:8] or "bin"
    key = f"{_ATTACH_PREFIX}/{expense_id}/{uuid.uuid4().hex}.{ext}"
    try:
        supabase.storage.from_(_BUCKET).upload(
            path=key,
            file=content,
            file_options={"content-type": content_type, "upsert": "true"},
        )
    except Exception as e:
        raise HTTPException(status_code=502, detail=f"Upload failed: {e}")

    record = {
        "expense_id": expense_id,
        "storage_path": f"storage://{_BUCKET}/{key}",
        "file_name": file.filename or f"attachment.{ext}",
        "kind": (kind or None),
        "uploaded_by": _user_id(request),
    }
    return supabase.table("expense_attachments").insert(record).execute().data[0]


# ─── Recurrentes ─────────────────────────────────────────────────────────────
@router.get("/recurring")
def list_recurring(request: Request, include_inactive: bool = Query(False)):
    q = _scoped(supabase.table("recurring_expenses").select(
        "*, vendors(name), departments(label), accounts(label)"
    ).order("description"), request)
    if not include_inactive:
        q = q.eq("active", True)
    rows = q.execute().data or []
    for r in rows:
        r["vendor_name"] = (r.pop("vendors", None) or {}).get("name")
        r["department_label"] = (r.pop("departments", None) or {}).get("label")
        r["account_label"] = (r.pop("accounts", None) or {}).get("label")
    return rows


@router.post("/recurring", status_code=201, dependencies=[Depends(require_edit("budget"))])
def create_recurring(data: RecurringCreate, request: Request):
    if not _valid_department(data.department_code):
        raise HTTPException(status_code=400, detail="Unknown or inactive department")
    _assert_can_edit(request, data.department_code)
    _check_account(data.account_code, "expense")
    if data.frequency not in _FREQUENCIES:
        raise HTTPException(status_code=400, detail="Invalid frequency")
    record = {k: v for k, v in data.model_dump().items() if v is not None}
    record["created_by"] = _user_id(request)
    return supabase.table("recurring_expenses").insert(record).execute().data[0]


def _covers_period(row: dict, year: int, month: int) -> bool:
    """¿La plantilla vence en este mes?

    monthly   → todos los meses del tramo
    quarterly → los meses del mes de inicio + 3, 6, 9
    annual    → solo el mes de inicio
    """
    start = row.get("start_date")
    if not start:
        return False
    s_year, s_month = int(start[0:4]), int(start[5:7])
    if (year, month) < (s_year, s_month):
        return False
    end = row.get("end_date")
    if end and (year, month) > (int(end[0:4]), int(end[5:7])):
        return False

    freq = row.get("frequency")
    if freq == "monthly":
        return True
    elapsed = (year - s_year) * 12 + (month - s_month)
    if freq == "quarterly":
        return elapsed % 3 == 0
    if freq == "annual":
        return elapsed % 12 == 0
    return False


@router.get("/recurring/pending")
def pending_recurring(request: Request, period: str = Query(..., description="'YYYY-MM'")):
    """Qué gastos recurrentes toca cargar este mes y todavía no se cargaron.

    NO crea nada: solo compara las plantillas activas contra los `expenses` que
    ya tienen ese (recurring_id, period). El alta la confirma el usuario en
    /recurring/generate — por eso el módulo no necesita cron y no puede inventar
    gastos que nadie revisó.
    """
    if not _PERIOD_RE.match(period):
        raise HTTPException(status_code=400, detail="period must be 'YYYY-MM'")
    return _pending_recurring(request, period)


def _pending_recurring(request: Request, period: str) -> list:
    """Lógica de GET /recurring/pending, sin el endpoint alrededor.

    `period` ya validado ('YYYY-MM'). La reusa la bandeja de pendientes de la
    campana (api/_lib/routers/notifications.py) para contar lo que falta cargar
    en el mes; pasa por el mismo _scoped(), así que cuenta solo lo que el caller
    vería en la pestaña.
    """
    year, month = int(period[0:4]), int(period[5:7])

    templates = _scoped(
        supabase.table("recurring_expenses")
        .select("*, vendors(name), departments(label), accounts(label)")
        .eq("active", True),
        request,
    ).execute().data or []
    due = [t for t in templates if _covers_period(t, year, month)]
    if not due:
        return []

    already = (
        supabase.table("expenses")
        .select("recurring_id")
        .eq("recurring_period", period)
        .in_("recurring_id", [t["id"] for t in due])
        .execute()
        .data
    ) or []
    done = {r["recurring_id"] for r in already}

    out = []
    for t in due:
        if t["id"] in done:
            continue
        out.append({
            "recurring_id": t["id"],
            "period": period,
            "description": t.get("description"),
            "amount": t.get("amount"),
            "department_code": t.get("department_code"),
            "department_label": (t.get("departments") or {}).get("label"),
            "account_code": t.get("account_code"),
            "account_label": (t.get("accounts") or {}).get("label"),
            "vendor_id": t.get("vendor_id"),
            "vendor_name": (t.get("vendors") or {}).get("name"),
        })
    return sorted(out, key=lambda r: (r["description"] or "").lower())


@router.post("/recurring/generate", status_code=201, dependencies=[Depends(require_edit("budget"))])
def generate_recurring(data: RecurringGenerate, request: Request):
    """Confirma en bloque los recurrentes de un período.

    Los crea en `draft` — son gastos esperados, no ejecutados. El monto se puede
    pisar por ítem si la factura vino distinta a la plantilla.
    """
    if not data.items:
        raise HTTPException(status_code=400, detail="Nothing to generate")

    created, skipped = [], []
    for item in data.items:
        if not _PERIOD_RE.match(item.period):
            raise HTTPException(status_code=400, detail=f"Invalid period: {item.period}")

        tpl = (
            supabase.table("recurring_expenses")
            .select("*")
            .eq("id", item.recurring_id)
            .execute()
            .data
        )
        if not tpl:
            raise HTTPException(status_code=404, detail=f"Recurring expense {item.recurring_id} not found")
        tpl = tpl[0]
        # Generar crea gastos: el permiso se valida por plantilla, no una vez
        # al entrar — la lista de ítems la manda el cliente.
        _assert_can_edit(request, tpl.get("department_code"))

        # El índice único (recurring_id, recurring_period) ya lo garantiza; se
        # chequea antes para poder reportar el salteo en vez de tirar un 409 y
        # perder los ítems que sí se podían crear.
        dup = (
            supabase.table("expenses")
            .select("id")
            .eq("recurring_id", item.recurring_id)
            .eq("recurring_period", item.period)
            .execute()
            .data
        )
        if dup:
            skipped.append({"recurring_id": item.recurring_id, "period": item.period})
            continue

        day = tpl.get("day_of_month") or 1
        expense_date = item.expense_date or f"{item.period}-{day:02d}"
        record = {
            "department_code": tpl["department_code"],
            "account_code": tpl["account_code"],
            "payee_type": "vendor" if tpl.get("vendor_id") else "other",
            "vendor_id": tpl.get("vendor_id"),
            "payee_name": None if tpl.get("vendor_id") else tpl.get("description"),
            "description": f"{tpl['description']} — {item.period}",
            "amount": item.amount if item.amount is not None else tpl["amount"],
            "expense_date": expense_date,
            "status": "draft",
            "recurring_id": tpl["id"],
            "recurring_period": item.period,
            "created_by": _user_id(request),
        }
        record = {k: v for k, v in record.items() if v is not None}
        row = supabase.table("expenses").insert(record).execute().data[0]
        created.append(row)
        # Es un alta de gasto como cualquier otra: entra al historial y les
        # avisa a los aprobadores del departamento.
        record_transition(
            "expense", row["id"], None, "draft", request,
            meta={"recurring_id": tpl["id"], "period": item.period},
        )
        _notify_expense_submitted(row, request)

    return {"created": created, "skipped": skipped}


# Las rutas con {recurring_id} van DESPUÉS de /pending y /generate: FastAPI
# resuelve en orden de declaración, así que un path param declarado antes se
# come las rutas literales que comparten prefijo.
@router.patch("/recurring/{recurring_id}", dependencies=[Depends(require_edit("budget"))])
def update_recurring(recurring_id: str, data: RecurringUpdate, request: Request):
    _row_or_403("recurring_expenses", recurring_id, request, "edit")
    updates = _patch_updates(data, "recurring_expenses")
    if "department_code" in updates:
        if not _valid_department(updates["department_code"]):
            raise HTTPException(status_code=400, detail="Unknown or inactive department")
        _assert_can_edit(request, updates["department_code"])
    if "account_code" in updates:
        _check_account(updates["account_code"], "expense")
    if "frequency" in updates and updates["frequency"] not in _FREQUENCIES:
        raise HTTPException(status_code=400, detail="Invalid frequency")
    updates["updated_at"] = _now_iso()
    result = supabase.table("recurring_expenses").update(updates).eq("id", recurring_id).execute()
    if not result.data:
        raise HTTPException(status_code=404, detail="Recurring expense not found")
    return result.data[0]


@router.delete("/recurring/{recurring_id}", dependencies=[Depends(require_edit("budget"))])
def delete_recurring(recurring_id: str, request: Request):
    """Baja lógica si ya generó gastos; borrado real si nunca generó ninguno."""
    _row_or_403("recurring_expenses", recurring_id, request, "edit")
    used = supabase.table("expenses").select("id").eq("recurring_id", recurring_id).limit(1).execute().data
    if used:
        result = supabase.table("recurring_expenses").update(
            {"active": False, "updated_at": _now_iso()}
        ).eq("id", recurring_id).execute()
        if not result.data:
            raise HTTPException(status_code=404, detail="Recurring expense not found")
        return {"ok": True, "deactivated": True}

    result = supabase.table("recurring_expenses").delete().eq("id", recurring_id).execute()
    if not result.data:
        raise HTTPException(status_code=404, detail="Recurring expense not found")
    return {"ok": True, "deactivated": False}


# ─── Ingresos ────────────────────────────────────────────────────────────────
@router.get("/revenues")
def list_revenues(
    request: Request,
    year: Optional[int] = Query(None),
    department: Optional[str] = Query(None),
    competition_id: Optional[str] = Query(None),
    status: Optional[str] = Query(None),
):
    q = _scoped(supabase.table("revenues").select(
        "*, competitions(name), departments(label), accounts(label)"
    ).order("expected_date", desc=True), request, department)
    if competition_id:
        q = q.eq("competition_id", competition_id)
    if status:
        q = q.eq("status", status)
    rows = q.execute().data or []
    for r in rows:
        r["competition_name"] = (r.pop("competitions", None) or {}).get("name")
        r["department_label"] = (r.pop("departments", None) or {}).get("label")
        r["account_label"] = (r.pop("accounts", None) or {}).get("label")

    if year:
        # Un ingreso cae en el año por la fecha que tenga: la de cobro si ya se
        # cobró, la esperada si todavía no.
        def _in_year(r):
            d = r.get("received_date") or r.get("expected_date") or ""
            return d[:4] == str(year)
        rows = [r for r in rows if _in_year(r)]
    return rows


@router.post("/revenues", status_code=201, dependencies=[Depends(require_edit("budget"))])
def create_revenue(data: RevenueCreate, request: Request):
    if not _valid_department(data.department_code):
        raise HTTPException(status_code=400, detail="Unknown or inactive department")
    _assert_can_edit(request, data.department_code)
    _check_account(data.account_code, "revenue")
    status = data.status or "expected"
    if status not in _REVENUE_STATUSES:
        raise HTTPException(status_code=400, detail="Invalid status")
    if status == "received" and not data.received_date:
        raise HTTPException(status_code=400, detail="A received date is required to mark revenue as received")
    if status != "received" and data.received_date:
        raise HTTPException(status_code=400, detail="Only received revenue can carry a received date")

    record = {k: v for k, v in data.model_dump().items() if v is not None}
    record["status"] = status
    record["created_by"] = _user_id(request)
    return supabase.table("revenues").insert(record).execute().data[0]


@router.patch("/revenues/{revenue_id}", dependencies=[Depends(require_edit("budget"))])
def update_revenue(revenue_id: str, data: RevenueUpdate, request: Request):
    current = _row_or_403("revenues", revenue_id, request, "edit")

    updates = _patch_updates(data, "revenues")
    if "department_code" in updates:
        if not _valid_department(updates["department_code"]):
            raise HTTPException(status_code=400, detail="Unknown or inactive department")
        _assert_can_edit(request, updates["department_code"])
    if "account_code" in updates:
        _check_account(updates["account_code"], "revenue")

    status = updates.get("status", current.get("status"))
    if status not in _REVENUE_STATUSES:
        raise HTTPException(status_code=400, detail="Invalid status")
    received_date = updates.get("received_date", current.get("received_date"))
    if status != "received":
        received_date = None
        updates["received_date"] = None
    if status == "received" and not received_date:
        raise HTTPException(status_code=400, detail="A received date is required to mark revenue as received")

    updates["updated_at"] = _now_iso()
    result = supabase.table("revenues").update(updates).eq("id", revenue_id).execute()
    if not result.data:
        raise HTTPException(status_code=404, detail="Revenue not found")
    return result.data[0]


@router.delete("/revenues/{revenue_id}", dependencies=[Depends(require_edit("budget"))])
def delete_revenue(revenue_id: str, request: Request):
    _row_or_403("revenues", revenue_id, request, "edit")
    result = supabase.table("revenues").delete().eq("id", revenue_id).execute()
    if not result.data:
        raise HTTPException(status_code=404, detail="Revenue not found")
    return {"ok": True}
